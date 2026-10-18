"""Tests for PPTX export."""

import tempfile
from pathlib import Path

import pytest

from colloquium.export import export_pptx

pytest.importorskip("pptx", reason="python-pptx not installed")

def test_export_scatter_chart_smoke():
    md_content = """---
title: Scatter
---

## XY Chart

```chart
type: scatter
data:
  datasets:
    - label: Series
      data:
        - {x: 1, y: 2}
        - {x: 2, y: 3}
        - {x: 3, y: 5}
```
"""

    with tempfile.TemporaryDirectory() as tmpdir:
        md_path = Path(tmpdir) / "scatter.md"
        md_path.write_text(md_content)

        result = export_pptx(str(md_path))

        assert Path(result).exists()
        assert Path(result).suffix == ".pptx"


class TestIframePrintSnapshots:
    def test_injects_snapshot_after_iframe_and_before_fallback(self, monkeypatch, tmp_path):
        from colloquium import export as export_module

        def fake_capture(browser, url, out_path):
            import pathlib
            pathlib.Path(out_path).write_bytes(b"\x89PNG" + b"0" * 32768)
            return True

        monkeypatch.setattr(export_module, "_capture_page_snapshot", fake_capture)
        html = (
            '<iframe class="colloquium-iframe" src="https://example.com/a?x=1&amp;y=2" '
            'width="100%" height="300" frameborder="0" style="border:none"></iframe>'
            '<div class="colloquium-iframe-print-fallback">card</div>'
        )
        out = export_module._inject_iframe_snapshots(html, browser="chrome")
        assert 'colloquium-iframe-print-snapshot' in out
        snap_at = out.index("colloquium-iframe-print-snapshot")
        assert out.index("</iframe>") < snap_at < out.index("colloquium-iframe-print-fallback")
        assert "data:image/png;base64," in out

    def test_failed_capture_leaves_markup_unchanged(self, monkeypatch):
        from colloquium import export as export_module

        monkeypatch.setattr(
            export_module, "_capture_page_snapshot", lambda *a, **k: False
        )
        html = (
            '<iframe class="colloquium-iframe" src="https://example.com/a" '
            'width="100%" height="300" frameborder="0" style="x"></iframe>'
            '<div class="colloquium-iframe-print-fallback">card</div>'
        )
        out = export_module._inject_iframe_snapshots(html, browser="chrome")
        assert out == html


def test_plain_export_honors_only_and_expands_toc(tmp_path):
    from pptx import Presentation

    md_path = tmp_path / "talk.md"
    md_path.write_text("""---
title: Plain
---

## Outline

```toc
```

---

<!-- only: html -->
<!-- layout: section-break -->

## Web only

---

<!-- layout: section-break -->

## Shared part

---

## Embeds

```iframe
src: https://example.com
```

```python
print("kept")
```
""")
    prs = Presentation(export_pptx(str(md_path)))
    texts = [[s.text_frame.text for s in slide.shapes if s.has_text_frame] for slide in prs.slides]

    assert [t[0] for t in texts] == ["Outline", "Shared part", "Embeds"]
    assert any("Shared part" in t for t in texts[0][1:])
    assert not any("Web only" in t for slide in texts for t in slide)
    embeds = "\n".join(texts[2])
    assert "example.com" not in embeds
    assert 'print("kept")' in embeds
