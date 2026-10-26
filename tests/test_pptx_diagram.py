"""Tests for HTML diagrams exported as native PPTX shapes."""

from pathlib import Path

import pytest

pytest.importorskip("pptx", reason="python-pptx not installed")

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from colloquium.export import _find_browser, _is_chromium_based, export_pptx
from colloquium.pptx_diagram import (
    DIAGRAM_ATTR,
    Diagram,
    _mark_diagrams,
    add_diagram,
    add_morph_transition,
    content_bounds,
    is_diagram_html,
    map_color,
)
from colloquium.pptx_template import TemplateConfig

_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
_MORPH = "{http://schemas.microsoft.com/office/powerpoint/2015/09/main}morph"

_browser = _find_browser()
needs_chromium = pytest.mark.skipif(
    _browser is None or not _is_chromium_based(_browser), reason="needs a Chromium-based browser"
)


def _cfg(tmp_path: Path) -> TemplateConfig:
    return TemplateConfig(
        template=tmp_path / "unused.pptx",
        layouts={},
        font="Inter",
        colors={"text": "280016", "muted": "706666", "accent": "FF0073", "link": "D10060"},
    )


def test_diagram_html_needs_nested_divs_or_svg():
    assert is_diagram_html('<div class="a"><div class="b">x</div></div>')
    assert is_diagram_html('<svg viewBox="0 0 10 10"><line x1="0" y1="0" x2="1" y2="1"/></svg>')
    # A legend of inline spans stays text, as do captions.
    assert not is_diagram_html('<div class="plot-legend"><span>USA</span></div>')
    assert not is_diagram_html('<p class="caption">Source</p>')
    assert not is_diagram_html("<!-- <div><div></div></div> --><p>x</p>")


def test_mark_diagrams_tags_each_root_in_source_order():
    content = """Intro

<!-- note -->
<div class="d1">
  <div>a</div>
</div>

<p class="caption">not a diagram</p>

<svg class="d2">
  <line x1="0" y1="0" x2="1" y2="1"/>
</svg>
"""
    marked, count = _mark_diagrams(content, 7)
    assert count == 2
    assert f'<div {DIAGRAM_ATTR}="7:0" class="d1">' in marked
    assert f'<svg {DIAGRAM_ATTR}="7:1" class="d2">' in marked
    assert '<p class="caption">' in marked


def test_map_color_follows_template_palette():
    pairs = [((0x4A, 0xA6, 0x91), (0xFF, 0x00, 0x73)), ((0x0B, 0x1A, 0x14), (0x28, 0x00, 0x16))]
    # Exact theme colors swap, alpha is not part of the match.
    assert map_color([0x4A, 0xA6, 0x91, 0.24], pairs) == (0xFF, 0x00, 0x73)
    assert map_color((0x0B, 0x1A, 0x14), pairs) == (0x28, 0x00, 0x16)
    # A tint of the accent becomes the same tint of the template accent.
    tint = map_color((0xB8, 0xD4, 0xC8), pairs)
    assert tint[0] == 255 and 120 < tint[1] < 170
    # A slightly lighter shade of the text color stays a lighter template text color.
    shade = map_color((0x24, 0x34, 0x2D), pairs)
    assert shade != (0x24, 0x34, 0x2D) and shade[0] > 0x28
    # White and unrelated colors are kept.
    assert map_color((255, 255, 255), pairs) == (255, 255, 255)
    assert map_color((0x22, 0x58, 0xC3), pairs) == (0x22, 0x58, 0xC3)


def _sample_diagram() -> Diagram:
    return Diagram(
        signature="div.d",
        x=60, y=180, w=600, h=300,
        theme={"--colloquium-accent": [74, 166, 145, 1], "--colloquium-heading": [11, 26, 20, 1]},
        items=[
            {"kind": "box", "key": "div.node#box", "order": 1, "x": 100, "y": 200, "w": 200, "h": 80,
             "radius": 12, "fill": [255, 255, 255, 1], "stroke": [74, 166, 145, 1], "stroke_width": 2},
            {"kind": "text", "key": "div.node>div.label#text", "order": 2, "x": 110, "y": 220, "w": 180, "h": 30,
             "align": "center", "wrap": True,
             "runs": [{"text": "Building models", "size": 24, "bold": True, "color": [11, 26, 20, 1]}]},
            {"kind": "box", "key": "div.tint#box", "order": 4, "z": 3, "x": 250, "y": 230, "w": 80, "h": 60,
             "radius": 0, "fill": [74, 166, 145, 0.24], "stroke": None, "stroke_width": 0,
             "runs": [{"text": "Overlap", "size": 19, "bold": True, "color": [11, 26, 20, 1]}],
             "padding": [0, 8, 0, 8], "align": "center", "valign": "middle"},
            {"kind": "line", "key": "svg>line:0", "order": 3, "points": [[200, 280], [200, 420]],
             "stroke": [74, 166, 145, 1], "stroke_width": 3, "dash": [10, 8], "arrow_end": True},
        ],
    )


def test_add_diagram_draws_named_native_shapes(tmp_path):
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_diagram(slide, _sample_diagram(), x0=1.0, y0=2.0, s=0.01, cfg=_cfg(tmp_path))

    shapes = list(slide.shapes)
    assert all(s.name.startswith("!!") for s in shapes)
    node, label, line, tint = shapes  # z-index 3 draws the tint last
    assert node.auto_shape_type is not None and node.left == Inches(1.4) and node.top == Inches(2.2)
    assert str(node.line.color.rgb) == "FF0073"  # theme accent -> template accent
    assert label.text_frame.text == "Building models"
    run = label.text_frame.paragraphs[0].runs[0]
    assert run.font.bold and run.font.name == "Inter" and str(run.font.color.rgb) == "280016"
    assert run.font.size.pt == pytest.approx(24 * 0.01 * 72, abs=0.5)
    assert line.shape_type == MSO_SHAPE_TYPE.LINE
    ln = line._element.spPr.find(qn("a:ln"))
    assert ln.find(qn("a:prstDash")).get("val") == "dash"
    assert ln.find(qn("a:tailEnd")).get("type") == "triangle"
    assert tint.text_frame.text == "Overlap"
    alpha = tint._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")).find(qn("a:alpha"))
    assert alpha.get("val") == "24000"


def test_morph_transition_sits_after_clr_map_with_a_fade_fallback():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_morph_transition(slide)
    add_morph_transition(slide)  # replaces, never stacks
    blocks = slide._element.findall(_MC)
    assert len(blocks) == 1
    assert blocks[0].getprevious().tag in (qn("p:clrMapOvr"), qn("p:cSld"))
    assert blocks[0].find(".//" + _MORPH).get("option") == "byObject"
    assert blocks[0].find(".//" + qn("p:fade")) is not None


def _write_deck(tmp_path: Path, body: str) -> Path:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.slides.add_slide(prs.slide_layouts[5])
    prs.save(tmp_path / "template.pptx")
    (tmp_path / "deck.css").write_text("""
.d { position: absolute; left: 60px; right: 60px; top: 180px; bottom: 60px; }
.d .node { position: absolute; width: 300px; padding: 16px; background: #fff;
           border: 2px solid var(--colloquium-accent); border-radius: 12px; text-align: center; }
.d .a { left: 100px; top: 40px; }
.step-2 .d .a { left: 600px; }
.d .b { left: 100px; top: 260px; }
.d .label { font-size: 24px; font-weight: 700; }
.fade-in { animation: fade 1s both; opacity: 0; }
@keyframes fade { to { opacity: 1; } }
""")
    md = tmp_path / "talk.md"
    md.write_text(f"""---
title: Talk
custom_css: deck.css
pptx:
  template: template.pptx
  layouts:
    title: 1
    content: 2
  colors: {{text: "280016", accent: "FF0073"}}
---

{body}""")
    return md


@needs_chromium
def test_diagram_steps_export_as_shapes_with_morph(tmp_path):
    step = """<div class="d">
  <svg viewBox="0 0 100 100" preserveAspectRatio="none" style="position:absolute;inset:0;width:100%;height:100%">
    <line x1="30" y1="20" x2="30" y2="80" stroke="black" stroke-width="2"/>
  </svg>
  <div class="node a"><div class="label">Building</div></div>
  <div class="node b fade-in"><div class="label">Serving</div></div>
</div>"""
    md = _write_deck(tmp_path, f"""<!-- class: step-1 -->
## Step one

{step}

---

<!-- class: step-2 -->
## Step two

{step}
""")
    prs = Presentation(export_pptx(str(md)))
    first, second = prs.slides
    names = [{s.name for s in slide.shapes if s.name.startswith("!!")} for slide in (first, second)]
    assert names[0] and names[0] == names[1]
    assert first._element.find(_MC) is None
    assert second._element.find(_MC) is not None

    def node(slide, text):
        label = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text == text)
        boxes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
        return label, boxes

    building_1, boxes_1 = node(first, "Building")
    building_2, _ = node(second, "Building")
    # Same diagram, same scale: the node moved right by its CSS offset only.
    assert building_2.left > building_1.left
    assert building_2.top == building_1.top
    # The finished fade-in counts: "Serving" is drawn, and the accent border is the template's.
    assert node(first, "Serving")[0] is not None
    assert {str(b.line.color.rgb) for b in boxes_1} == {"FF0073"}
    assert any(s.shape_type == MSO_SHAPE_TYPE.LINE for s in first.shapes)
    # The snapshot page is cleaned up.
    assert not list(tmp_path.glob(".colloquium-pptx-*"))


def test_content_bounds_cover_what_is_drawn_not_the_root_box():
    diagram = Diagram(signature="div.canvas", x=60, y=180, w=1160, h=490, items=[
        {"kind": "box", "key": "a", "x": 450, "y": 200, "w": 380, "h": 120},
        {"kind": "line", "key": "l", "points": [[640, 320], [640, 520]], "stroke_width": 2},
        {"kind": "text", "key": "t", "x": 470, "y": 520, "w": 340, "h": 40},
    ], theme={})
    assert content_bounds(diagram) == (450, 200, 380, 360)
    empty = Diagram(signature="div.canvas", x=60, y=180, w=1160, h=490, items=[], theme={})
    assert content_bounds(empty) == (60, 180, 1160, 490)
