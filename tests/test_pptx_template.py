"""Tests for template-aware PPTX export."""

from pathlib import Path

import pytest

pytest.importorskip("pptx", reason="python-pptx not installed")

from pptx import Presentation
from pptx.oxml.ns import qn

from colloquium.export import export_pptx


def _make_template(path: Path) -> Path:
    """A 3-slide template built from python-pptx's stock layouts:
    1 = Title Slide, 2 = Section Header, 3 = Title Only."""
    prs = Presentation()
    for layout_idx in (0, 2, 5):
        prs.slides.add_slide(prs.slide_layouts[layout_idx])
    prs.save(path)
    return path


def _write_deck(tmp_path: Path, body: str, pptx_block: str | None = None) -> Path:
    _make_template(tmp_path / "template.pptx")
    if pptx_block is None:
        pptx_block = """pptx:
  template: template.pptx
  layouts:
    title: {slide: 1, subtitle: 1}
    section-break: {slide: 2, eyebrow: 1}
    content: 3
  font: Inter
  colors: {text: "280016", accent: "FF0073", link: "D10060"}
"""
    md = tmp_path / "talk.md"
    md.write_text(f"---\ntitle: Talk\n{pptx_block}footer:\n  right: \"Me {{n}}/{{N}}\"\n---\n\n{body}")
    return md


def _texts(slide) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def test_builds_on_template_layouts_and_drops_sample_slides(tmp_path):
    md = _write_deck(tmp_path, """# My Talk

Subtitle line

---

<!-- layout: section-break -->
<!-- eyebrow: Part 1 -->

## The Section

---

## Content

Hello **world** and [a link](https://example.com).
""")
    out = export_pptx(str(md))
    prs = Presentation(out)

    assert len(prs.slides) == 3
    names = [s.slide_layout.name for s in prs.slides]
    assert names == ["Title Slide", "Section Header", "Title Only"]
    # Layouts nothing uses are pruned.
    assert {layout.name for layout in prs.slide_masters[0].slide_layouts} == set(names)

    title, section, content = prs.slides
    assert "My Talk" in _texts(title)
    assert "Subtitle line" in _texts(title)
    assert _texts(section)[:2] == ["The Section", "Part 1"]

    body = next(s for s in content.shapes if s.has_text_frame and "Hello" in s.text_frame.text)
    runs = body.text_frame.paragraphs[0].runs
    bold = next(r for r in runs if r.text == "world")
    assert bold.font.bold
    link = next(r for r in runs if r.text == "a link")
    assert link.hyperlink.address == "https://example.com"
    assert str(link.font.color.rgb) == "D10060"
    assert all(r.font.name == "Inter" for r in runs)
    # Footer only on content slides, with {n}/{N} filled in.
    assert "Me 3/3" in _texts(content)
    assert not any("Me 1/3" in t for t in _texts(title))


def test_unused_placeholders_are_removed(tmp_path):
    md = _write_deck(tmp_path, """<!-- layout: section-break -->

## Only A Title
""")
    prs = Presentation(export_pptx(str(md)))
    slide = prs.slides[0]
    assert [ph.placeholder_format.idx for ph in slide.placeholders] == [0]


def test_lists_get_native_bullets_and_numbering(tmp_path):
    md = _write_deck(tmp_path, """## Lists

- alpha
- beta

5. fifth
6. sixth
""")
    prs = Presentation(export_pptx(str(md)))
    paras = [
        p for shape in prs.slides[0].shapes if shape.has_text_frame
        for p in shape.text_frame.paragraphs if p.text in {"alpha", "fifth", "sixth"}
    ]
    alpha, fifth, sixth = paras
    assert alpha._p.pPr.find(qn("a:buChar")) is not None
    auto = fifth._p.pPr.find(qn("a:buAutoNum"))
    assert auto.get("type") == "arabicPeriod" and auto.get("startAt") == "5"
    assert sixth._p.pPr.find(qn("a:buAutoNum")).get("startAt") == "6"


def test_stacked_h2_lines_merge_into_title(tmp_path):
    md = _write_deck(tmp_path, """## First line
## Second line

Body text
""")
    prs = Presentation(export_pptx(str(md)))
    title = prs.slides[0].placeholders[0]
    assert title.text_frame.text == "First line\vSecond line"
    assert any("Body text" in t for t in _texts(prs.slides[0]))


def test_image_with_html_legend_and_columns(tmp_path):
    from PIL import Image

    Image.new("RGB", (400, 200), "white").save(tmp_path / "plot.png")
    md = _write_deck(tmp_path, """<!-- columns: 50/50 -->
## Figure

Left text

|||

![plot](plot.png)

<div class="plot-legend">
  <span class="plot-legend-title">Region</span>
  <span class="plot-legend-item"><span class="plot-legend-line" style="--legend-color: #2258C3;"></span>USA</span>
</div>
""")
    prs = Presentation(export_pptx(str(md)))
    slide = prs.slides[0]
    pictures = [s for s in slide.shapes if s.shape_type == 13]
    assert len(pictures) == 1
    left_text = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "Left text")
    # Second column starts right of the first.
    assert pictures[0].left > left_text.left + left_text.width
    legend = next(s for s in slide.shapes if s.has_text_frame and "USA" in s.text_frame.text)
    swatch = next(r for r in legend.text_frame.paragraphs[0].runs if "━" in r.text)
    assert str(swatch.font.color.rgb) == "2258C3"


def test_cli_template_override_and_bad_config(tmp_path):
    md = _write_deck(tmp_path, "## Hi\n\ntext\n", pptx_block="")
    template = str(tmp_path / "template.pptx")
    with pytest.raises(ValueError, match="content"):
        export_pptx(str(md), template=template)

    md = _write_deck(tmp_path, "## Hi\n", pptx_block="pptx:\n  layouts: {content: 9}\n")
    with pytest.raises(ValueError, match="template has 3 slides"):
        export_pptx(str(md), template=template)

    md = _write_deck(tmp_path, "## Hi\n", pptx_block="pptx:\n  layouts: {content: 3}\n")
    with pytest.raises(FileNotFoundError):
        export_pptx(str(md), template=str(tmp_path / "missing.pptx"))
    out = export_pptx(str(md), template=template)
    assert len(Presentation(out).slides) == 1


def test_citations_and_references_slide(tmp_path):
    (tmp_path / "refs.bib").write_text(
        "@misc{doe2026,\n  title = {A Paper},\n  author = {Doe, Jane and Roe, Rick},\n"
        "  year = {2026},\n  url = {https://example.com/paper}\n}\n"
    )
    md = _write_deck(tmp_path, """<!-- cite: doe2026 -->
## Cited

Claim
""")
    text = md.read_text().replace("title: Talk\n", "title: Talk\nbibliography: refs.bib\n")
    md.write_text(text)
    prs = Presentation(export_pptx(str(md)))
    assert len(prs.slides) == 2
    assert any("Doe & Roe, 2026" in t for t in _texts(prs.slides[0]))
    refs = prs.slides[1]
    assert refs.placeholders[0].text_frame.text == "References"
    assert any("Doe, J., and Roe, R. “A Paper.” 2026." in t for t in _texts(refs))


def test_post_reference_slides_follow_html_numbering(tmp_path):
    (tmp_path / "refs.bib").write_text(
        "@misc{a2026,\n  title = {A},\n  author = {Ay, Ann},\n  year = {2026}\n}\n"
        "@misc{b2026,\n  title = {B},\n  author = {Bee, Bob},\n  year = {2026}\n}\n"
    )
    md = _write_deck(tmp_path, """<!-- after: references -->
## Appendix

Backup [@b2026]

---

## Main

Claim [@a2026]
""")
    text = md.read_text().replace(
        "title: Talk\n", "title: Talk\nbibliography: refs.bib\ncitation_style: numeric\n"
    )
    md.write_text(text)
    prs = Presentation(export_pptx(str(md)))
    main, refs, appendix = prs.slides
    assert main.placeholders[0].text_frame.text == "Main"
    assert appendix.placeholders[0].text_frame.text == "Appendix"
    # The main slide is rendered first, so it gets [1] even though the
    # appendix comes first in the source.
    assert any("Claim [1]" in t for t in _texts(main))
    assert any("Backup [2]" in t for t in _texts(appendix))
    # {N} excludes post-reference slides, matching the HTML footer.
    assert "Me 1/2" in _texts(main)
    assert "Me 2/2" in _texts(refs)
    assert "Me 3/2" in _texts(appendix)


def test_only_directive_picks_slides_per_format(tmp_path):
    md = _write_deck(tmp_path, """<!-- only: html -->
## Interactive demo

Live iframe

---

<!-- only: pptx -->
## Static demo

Screenshot instead

---

## Shared

Both
""")
    prs = Presentation(export_pptx(str(md)))

    assert [s.placeholders[0].text_frame.text for s in prs.slides] == ["Static demo", "Shared"]
    assert "Me 2/2" in _texts(prs.slides[1])


def test_toc_block_becomes_a_chapter_list(tmp_path):
    md = _write_deck(tmp_path, """## Outline

```toc
```

---

<!-- layout: section-break -->

## Why open?

---

<!-- chapter: Adoption -->
## Agenda

```toc
```
""")
    prs = Presentation(export_pptx(str(md)))
    outline, _, agenda = prs.slides

    assert any("Why open?\nAdoption" in t for t in _texts(outline))
    body = next(s for s in agenda.shapes if s.has_text_frame and "Adoption" in s.text_frame.text)
    current = next(r for p in body.text_frame.paragraphs for r in p.runs if r.text == "Adoption")
    assert current.font.bold
    numbering = body.text_frame.paragraphs[0]._p.find(f"{qn('a:pPr')}/{qn('a:buAutoNum')}")
    assert numbering is not None


def test_element_blocks_are_skipped_not_printed(tmp_path, capsys):
    md = _write_deck(tmp_path, """## Embeds

Intro

```iframe
src: https://example.com
```

```builtwith
repo: natolambert/colloquium
```
""")
    prs = Presentation(export_pptx(str(md)))

    text = "\n".join(_texts(prs.slides[0]))
    assert "Intro" in text
    assert "example.com" not in text and "natolambert" not in text
    err = capsys.readouterr().err
    assert "```iframe blocks are not supported" in err
    assert "```builtwith blocks are not supported" in err



def test_html_img_tags_become_pictures_with_centered_captions(tmp_path):
    from PIL import Image

    Image.new("RGB", (400, 200), "white").save(tmp_path / "chart.png")
    md = _write_deck(tmp_path, """## Chart

<img src="chart.png" alt="A chart">
<p class="chart-caption"><a href="https://example.com">Source, 2026</a>: what it shows</p>
""")
    prs = Presentation(export_pptx(str(md)))
    slide = prs.slides[0]
    assert sum(1 for sh in slide.shapes if sh.shape_type == 13) == 1
    caption = next(sh for sh in slide.shapes if sh.has_text_frame and "Source, 2026" in sh.text_frame.text)
    from pptx.enum.text import PP_ALIGN

    assert caption.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    # Source spacing and the link survive: no space before the colon.
    assert caption.text_frame.text == "Source, 2026: what it shows"
    runs = caption.text_frame.paragraphs[0].runs
    assert runs[0].hyperlink.address == "https://example.com" and runs[-1].hyperlink.address is None
    # The caption box is centered on the figure, so it shares the figure's center.
    picture = next(sh for sh in slide.shapes if sh.shape_type == 13)
    assert abs((caption.left + caption.width / 2) - (picture.left + picture.width / 2)) < 20000
    assert caption.top > picture.top + picture.height


def _svg_plot(path: Path, title: str = "Old title") -> None:
    path.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="400" height="200" viewBox="0 0 400 200" aria-labelledby="t">'
        f'<title id="t">{title}</title><rect width="400" height="200" fill="#eee"/>'
        f'<text data-role="title" x="10" y="30">{title}</text></svg>'
    )


def test_svg_size_reads_width_height_or_viewbox(tmp_path):
    from colloquium.pptx_template import _svg_size

    _svg_plot(tmp_path / "plot.svg")
    assert _svg_size(tmp_path / "plot.svg") == (400.0, 200.0)
    (tmp_path / "vb.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 800 300"></svg>')
    assert _svg_size(tmp_path / "vb.svg") == (800.0, 300.0)


def test_retitle_svg_swaps_plot_title():
    from colloquium.build import retitle_svg

    svg = ('<svg xmlns="http://www.w3.org/2000/svg" aria-labelledby="t"><title id="t">Old</title>'
           '<text data-role="title">Old</text><text data-role="title">wrapped</text></svg>')
    out = retitle_svg(svg, "New heading")
    assert out.count("New heading") == 2 and "Old" not in out and "wrapped" not in out
    assert retitle_svg("<svg><title>x</title></svg>", "y") is None


def test_svg_figures_are_rasterized_or_skipped_with_a_warning(tmp_path, capsys):
    from colloquium.export import _find_browser, _is_chromium_based

    _svg_plot(tmp_path / "plot.svg")
    md = _write_deck(tmp_path, """<!-- figure-title: heading -->
## New heading

![Plot](plot.svg)
""")
    prs = Presentation(export_pptx(str(md)))
    pictures = [sh for sh in prs.slides[0].shapes if sh.shape_type == 13]
    # The plot carries the heading in its own title, so the slide doesn't repeat it.
    assert not any("New heading" in t for t in _texts(prs.slides[0]))
    browser = _find_browser()
    if browser and _is_chromium_based(browser):
        assert len(pictures) == 1 and pictures[0].image.ext == "png"
        assert pictures[0].image.size[0] >= 800  # rendered at 2x
    else:
        assert not pictures
        assert "rasterize SVG; skipped" in capsys.readouterr().err


def test_headings_shrink_to_stay_on_one_line_or_fit_their_box():
    from colloquium.pptx_template import (
        _HEADING_CHAR_EM,
        _HEADING_MIN_FRACTION,
        _estimate_lines,
        _fit_heading_size,
    )

    short = ["A short heading"]
    assert _fit_heading_size(short, 31, 8.9) == 31
    long = ["Open weights trail the frontier by a few months now"]
    size = _fit_heading_size(long, 31, 8.9)
    assert size < 31 and _estimate_lines(long[0], size, 8.9, _HEADING_CHAR_EM) == 1
    # Too long for one line even at the floor: the largest size that wraps only once.
    very_long = ["Revenue growth is key to continued training-scaling, and often captured elsewhere"]
    size = _fit_heading_size(very_long, 31, 8.9)
    assert size > 31 * _HEADING_MIN_FRACTION and _estimate_lines(very_long[0], size, 8.9, _HEADING_CHAR_EM) == 2
    huge = ["x " * 200]
    assert _fit_heading_size(huge, 31, 8.9) == pytest.approx(31 * _HEADING_MIN_FRACTION, abs=0.5)
    # Full-bleed headings (section breaks) shrink until they fit the given height.
    section = ["The task for open models -- build the cheapest, most adaptable, private intelligence"]
    fitted = _fit_heading_size(section, 54, 8.44, max_h=1.6)
    assert fitted < 54
    assert _estimate_lines(section[0], fitted, 8.44, _HEADING_CHAR_EM) * fitted * 1.15 / 72 <= 1.6
