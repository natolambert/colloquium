"""Template-aware PPTX export.

Builds native, editable slides on the layouts of an existing ``.pptx``
template (e.g. a conference's speaker template), so backgrounds, logos, and
heading styles come from the template while text, lists, links, and figures
come from the markdown. Configured with a ``pptx:`` frontmatter block::

    pptx:
      template: event-template.pptx   # relative to the markdown file
      layouts:                        # colloquium layout -> template slide number
        title: {slide: 4, eyebrow: 3, subtitle: 2, photo: 4, name: 6, affiliation: 5}
        section-break: {slide: 9, eyebrow: 2}
        content: {slide: 22, eyebrow: 3}
      font: Inter                     # body text font (headings use the template's)
      colors: {text: "280016", muted: "706666", accent: "FF0073", link: "D10060"}
      image_background: "FFFFFF"      # optional rounded card behind figures

Template slide numbers are 1-based, as shown in PowerPoint/Keynote. The
numbers after each role (``eyebrow: 3``) are placeholder ``idx`` values on
that slide's layout; the heading is ``idx`` 0 unless ``heading:`` says
otherwise. Placeholders that receive no content are removed, and layouts no
slide uses are pruned so the output doesn't carry the template's sample media.
"""

from __future__ import annotations

import html as html_module
import io
import re
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from colloquium.md import create_base_md
from colloquium.elements import BLOCK_LANGUAGES
from colloquium.pptx_diagram import (
    Diagram,
    add_diagram,
    add_morph_transition,
    content_bounds,
    is_diagram_html,
    snapshot_diagrams,
)

_ROLES = ("heading", "eyebrow", "subtitle", "name", "affiliation", "photo")

# Average glyph advance as a fraction of the font size. Deliberately a little
# generous so estimated text heights err toward fitting.
_BODY_CHAR_EM = 0.53
_HEADING_CHAR_EM = 0.5
_LINE_HEIGHT = 1.2
_MIN_SCALE = 0.6
# Content headings shrink this far below the template's size to stay on one line
# per `##` (a figure under them gets the height back).
_HEADING_MIN_FRACTION = 0.72
# Padding of the card drawn behind figures when pptx.image_background is set.
_CARD_PAD = 0.08
# Full-slide diagrams may grow past their HTML size by this factor to fill the body.
_DIAGRAM_MAX_ZOOM = 1.25

_COLUMN_SPLIT_RE = re.compile(r"^\s*\|\|\|\s*$", re.MULTILINE)
_HTML_COMMENT_RE = re.compile(r"<!--.*?-->", re.DOTALL)
_CITATION_RE = re.compile(r"\[@([^\]]+)\]")
_MD_IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)[^)]*\)")
_HTML_IMG_RE = re.compile(r"<img\b[^>]*\bsrc=[\"']([^\"']+)[\"'][^>]*>", re.IGNORECASE)
_HTML_CAPTION_CLASS_RE = re.compile(r"class=[\"'][^\"']*caption[^\"']*[\"']", re.IGNORECASE)
_TITLE_COMPONENT_RES = {
    "eyebrow": re.compile(r'<div class="colloquium-title-eyebrow">(.*?)</div>', re.DOTALL),
    "name": re.compile(r'<p class="colloquium-title-name">(.*?)</p>', re.DOTALL),
    "note": re.compile(r'<p class="colloquium-title-note">(.*?)</p>', re.DOTALL),
}
_TITLE_META_RE = re.compile(r'<div class="colloquium-title-meta">(.*?)</div>', re.DOTALL)
_TITLE_RULE_RE = re.compile(r'<div class="colloquium-title-rule">\s*</div>')
_HTML_P_RE = re.compile(r"<p[^>]*>(.*?)</p>", re.DOTALL)
_TAG_RE = re.compile(r"<[^>]+>")
_HEX_RE = re.compile(r"#([0-9a-fA-F]{6}|[0-9a-fA-F]{3})\b")


# ===== Config =====


@dataclass
class LayoutSpec:
    slide: int
    roles: dict[str, int] = field(default_factory=dict)


@dataclass
class TemplateConfig:
    template: Path
    layouts: dict[str, LayoutSpec]
    font: str | None = None
    colors: dict[str, str] = field(default_factory=dict)
    image_background: str | None = None


def _hex(value, default: str | None = None) -> str | None:
    if value is None:
        return default
    value = str(value).strip().lstrip("#")
    if len(value) == 3:
        value = "".join(c * 2 for c in value)
    if not re.fullmatch(r"[0-9a-fA-F]{6}", value):
        raise ValueError(f"pptx: invalid color {value!r} (expected hex like '280016')")
    return value.upper()


def parse_template_config(raw: dict, base_dir: Path, template_override: str | None = None) -> TemplateConfig:
    """Validate the ``pptx:`` frontmatter block."""
    raw = dict(raw or {})
    template = template_override or raw.get("template")
    if not template:
        raise ValueError("pptx: no template given (set pptx.template or pass --template)")
    template_path = Path(template).expanduser()
    if not template_path.is_absolute():
        template_path = (Path.cwd() if template_override else base_dir) / template_path
    if not template_path.is_file():
        raise FileNotFoundError(f"pptx: template not found: {template_path}")

    layouts: dict[str, LayoutSpec] = {}
    for name, spec in (raw.get("layouts") or {}).items():
        if isinstance(spec, int):
            layouts[str(name)] = LayoutSpec(slide=spec)
            continue
        if not isinstance(spec, dict) or not isinstance(spec.get("slide"), int):
            raise ValueError(f"pptx.layouts.{name}: expected a slide number or {{slide: N, ...}}")
        roles = {}
        for role, idx in spec.items():
            if role == "slide":
                continue
            if role not in _ROLES:
                raise ValueError(f"pptx.layouts.{name}: unknown role {role!r} (known: {', '.join(_ROLES)})")
            roles[role] = int(idx)
        layouts[str(name)] = LayoutSpec(slide=spec["slide"], roles=roles)
    if "content" not in layouts:
        raise ValueError("pptx.layouts needs at least a 'content' entry")

    colors_raw = raw.get("colors") or {}
    colors = {
        "text": _hex(colors_raw.get("text"), "1A1A2E"),
        "muted": _hex(colors_raw.get("muted"), "6B7280"),
        "accent": _hex(colors_raw.get("accent"), "0F3460"),
    }
    colors["link"] = _hex(colors_raw.get("link"), colors["accent"])

    return TemplateConfig(
        template=template_path,
        layouts=layouts,
        font=raw.get("font"),
        colors=colors,
        image_background=_hex(raw.get("image_background")),
    )


# ===== Markdown -> block model =====


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    href: str | None = None
    color: str | None = None
    line_break: bool = False


@dataclass
class Para:
    runs: list[Run]
    kind: str = "p"  # p | li | h | small
    ordered: bool = False
    number: int = 1
    level: int = 0
    align: str = "left"

    @property
    def text(self) -> str:
        return "".join("\n" if r.line_break else r.text for r in self.runs)


@dataclass
class TextBlock:
    paras: list[Para]


@dataclass
class ImageBlock:
    path: Path
    legend: TextBlock | None = None
    heading: str | None = None  # figure-title: heading swaps this into an SVG plot's title


@dataclass
class DiagramBlock:
    """An HTML diagram, drawn as native shapes from the browser's layout (see pptx_diagram)."""

    diagram: Diagram
    max_scale: float  # inches per CSS px at the HTML's own size; diagrams are never enlarged past it
    legend: TextBlock | None = None


class _SwatchHTMLParser(HTMLParser):
    """Flatten an HTML block to runs. Empty elements carrying a color in their
    inline style (legend swatches, markers, lines) become colored glyphs."""

    _VOID = {"br", "img", "hr", "input", "meta", "link"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.runs: list[Run] = []
        self._stack: list[tuple[str, str | None, str, int]] = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "br":
            self.runs.append(Run("", line_break=True))
            return
        if tag in self._VOID:
            return
        match = _HEX_RE.search(attrs.get("style") or "")
        color = _hex(match.group(1)) if match else None
        self._stack.append((tag, color, attrs.get("class") or "", len(self.runs), attrs.get("href")))

    def handle_endtag(self, tag):
        if not self._stack:
            return
        open_tag, color, cls, start, _ = self._stack.pop()
        if color and len(self.runs) == start:
            glyph = "━" if "line" in cls else "●" if "marker" in cls and "square" not in cls else "■"
            if self.runs and not self.runs[-1].text.endswith(" "):
                self.runs.append(Run("  "))
            self.runs.append(Run(glyph + " ", color=color))
        elif open_tag in {"div", "p", "li"} and self.runs and not self.runs[-1].text.endswith(" "):
            self.runs.append(Run(" "))

    def handle_data(self, data):
        # Keep the source's spacing: `<a>Link</a>: text` stays "Link: text".
        text = re.sub(r"\s+", " ", data)
        if not text:
            return
        if text.startswith(" ") and self.runs and not self.runs[-1].text.endswith(" ") and not self.runs[-1].line_break:
            self.runs.append(Run(" "))
        if text.strip():
            href = next((entry[4] for entry in reversed(self._stack) if entry[4]), None)
            self.runs.append(Run(text.strip(), href=href))
            if text.endswith(" "):
                self.runs.append(Run(" "))


def _html_to_runs(html: str) -> list[Run]:
    parser = _SwatchHTMLParser()
    parser.feed(html)
    parser.close()
    runs = parser.runs
    while runs and not runs[-1].text.strip() and not runs[-1].line_break:
        runs.pop()
    while runs and not runs[0].text.strip():
        runs.pop(0)
    return runs


def _inline_runs(children) -> list[Run]:
    runs: list[Run] = []
    bold = italic = False
    href = None
    for tok in children or []:
        t = tok.type
        if t == "text":
            runs.append(Run(tok.content, bold=bold, italic=italic, href=href))
        elif t == "code_inline":
            runs.append(Run(tok.content, bold=bold, italic=italic, code=True, href=href))
        elif t == "softbreak":
            runs.append(Run(" ", bold=bold, italic=italic, href=href))
        elif t == "hardbreak":
            runs.append(Run("", line_break=True))
        elif t == "strong_open":
            bold = True
        elif t == "strong_close":
            bold = False
        elif t == "em_open":
            italic = True
        elif t == "em_close":
            italic = False
        elif t == "link_open":
            href = tok.attrGet("href")
        elif t == "link_close":
            href = None
        elif t == "html_inline":
            if re.match(r"<br\s*/?>", tok.content, re.IGNORECASE):
                runs.append(Run("", line_break=True))
        elif t == "image":
            continue
        elif tok.content:
            runs.append(Run(tok.content, bold=bold, italic=italic, href=href))
    return runs


def _parse_blocks(
    markdown: str, base_dir: Path, warnings: list[str], next_diagram=None,
) -> list[TextBlock | ImageBlock | DiagramBlock]:
    """Convert a slide region's markdown into text, image, and diagram blocks.

    ``next_diagram`` returns the slide's next laid-out diagram as a DiagramBlock
    (or None when there is no layout for it), called once per diagram HTML
    block in source order.
    """
    tokens = create_base_md().parse(markdown)
    blocks: list[TextBlock | ImageBlock] = []
    current: list[Para] = []

    def flush():
        if current:
            blocks.append(TextBlock(list(current)))
            current.clear()

    list_stack: list[dict] = []
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        t = tok.type
        if t in ("bullet_list_open", "ordered_list_open"):
            start = int(tok.attrGet("start") or 1)
            list_stack.append({"ordered": t == "ordered_list_open", "n": start})
        elif t in ("bullet_list_close", "ordered_list_close"):
            list_stack.pop()
        elif t == "list_item_open":
            pass
        elif t == "paragraph_open":
            inline = tokens[i + 1]
            images = [c for c in inline.children or [] if c.type == "image"]
            runs = _inline_runs(inline.children)
            if images and not "".join(r.text for r in runs).strip() and not list_stack:
                flush()
                for img in images:
                    src = img.attrGet("src") or ""
                    path = _resolve_local(src, base_dir)
                    if path is None:
                        warnings.append(f"image not found or remote, skipped: {src}")
                        continue
                    blocks.append(ImageBlock(path))
            elif runs:
                if list_stack:
                    ctx = list_stack[-1]
                    current.append(Para(runs, kind="li", ordered=ctx["ordered"],
                                        number=ctx["n"], level=len(list_stack) - 1))
                    ctx["n"] += 1
                else:
                    current.append(Para(runs))
            i += 2
        elif t == "heading_open":
            current.append(Para(_inline_runs(tokens[i + 1].children), kind="h"))
            i += 2
        elif t == "html_block" and next_diagram is not None and is_diagram_html(tok.content) and (
            diagram_block := next_diagram()
        ) is not None:
            flush()
            if diagram_block.diagram.items:
                blocks.append(diagram_block)
        elif t == "html_block":
            html = _HTML_COMMENT_RE.sub("", tok.content).strip()
            # <img> tags are figures like markdown images; a <p class="…caption"> (or any
            # text that follows an <img>) is centered under the figure.
            img_srcs = _HTML_IMG_RE.findall(html)
            for src in img_srcs:
                path = _resolve_local(html_module.unescape(src), base_dir)
                if path is None:
                    warnings.append(f"image not found or remote, skipped: {src}")
                    continue
                flush()
                blocks.append(ImageBlock(path))
            html = _HTML_IMG_RE.sub("", html).strip()
            if html:
                runs = _html_to_runs(html)
                if runs:
                    caption = bool(img_srcs) or bool(_HTML_CAPTION_CLASS_RE.search(html))
                    legend = TextBlock([Para(runs, kind="small", align="center" if caption else "right")])
                    if blocks and isinstance(blocks[-1], ImageBlock) and not current and blocks[-1].legend is None:
                        blocks[-1].legend = legend
                    else:
                        flush()
                        blocks.append(legend)
        elif t in ("fence", "code_block"):
            if tok.info.strip() in BLOCK_LANGUAGES:
                warnings.append(f"```{tok.info.strip()} blocks are not supported in template export; skipped")
            else:
                for line in tok.content.rstrip("\n").splitlines():
                    current.append(Para([Run(line or " ", code=True)], kind="small"))
        elif t == "table_open":
            j = i
            rows: list[list[Run]] = []
            row: list[Run] = []
            while tokens[j].type != "table_close":
                if tokens[j].type == "inline":
                    if row:
                        row.append(Run("  |  "))
                    row.extend(_inline_runs(tokens[j].children))
                elif tokens[j].type == "tr_close":
                    rows.append(row)
                    row = []
                j += 1
            for r_i, r in enumerate(rows):
                if r_i == 0:
                    for run in r:
                        run.bold = True
                current.append(Para(r))
            i = j
        elif t == "blockquote_open":
            pass
        elif t == "hr":
            flush()
        i += 1
    flush()
    return blocks


def _resolve_local(src: str, base_dir: Path) -> Path | None:
    if re.match(r"^[a-z][a-z0-9+.-]*:", src, re.IGNORECASE):
        return None
    path = base_dir / src
    return path if path.is_file() else None


# ===== Measurement =====


def _estimate_lines(text: str, size_pt: float, width_in: float, char_em: float) -> int:
    chars_per_line = max(1.0, (width_in * 72) / (size_pt * char_em))
    lines = 0
    for segment in text.split("\n"):
        words = segment.split()
        if not words:
            lines += 1
            continue
        line_len = 0
        seg_lines = 1
        for word in words:
            add = len(word) + (1 if line_len else 0)
            if line_len + add > chars_per_line and line_len:
                seg_lines += 1
                line_len = len(word)
            else:
                line_len += add
        lines += seg_lines
    return lines


@dataclass
class _Metrics:
    base_pt: float

    def size(self, para: Para, scale: float) -> float:
        factor = {"h": 1.1, "small": 0.62}.get(para.kind, 1.0)
        return self.base_pt * scale * factor

    def indent_in(self, para: Para, scale: float) -> float:
        if para.kind != "li":
            return 0.0
        return (0.32 + 0.28 * para.level) * scale * self.base_pt / 13.5

    def space_before(self, para: Para, prev: Para | None, scale: float) -> float:
        if prev is None:
            return 0.0
        size = self.size(para, scale)
        if para.kind == "li" and prev.kind == "li":
            return size * 0.3
        return size * 0.55

    def text_height(self, block: TextBlock, width_in: float, scale: float) -> float:
        total = 0.0
        prev = None
        for para in block.paras:
            size = self.size(para, scale)
            width = width_in - self.indent_in(para, scale)
            lines = _estimate_lines(para.text, size, width, _BODY_CHAR_EM)
            total += self.space_before(para, prev, scale) + lines * size * _LINE_HEIGHT
            prev = para
        return total / 72


# ===== Region layout =====


@dataclass
class _Placed:
    kind: str  # text | image | diagram
    block: TextBlock | ImageBlock | DiagramBlock
    x: float
    y: float
    w: float
    h: float


def _svg_size(path: Path) -> tuple[float, float]:
    """Intrinsic size of an SVG in CSS px from its width/height or viewBox (1600x900 if unknown)."""
    import xml.etree.ElementTree as ET

    try:
        root = ET.parse(path).getroot()
    except ET.ParseError:
        return (1600.0, 900.0)

    def px(value: str | None) -> float | None:
        m = re.match(r"\s*([0-9.]+)", value or "")
        return float(m.group(1)) if m and float(m.group(1)) > 0 else None

    width, height = px(root.get("width")), px(root.get("height"))
    view_box = (root.get("viewBox") or "").replace(",", " ").split()
    if (width is None or height is None) and len(view_box) == 4:
        try:
            width, height = float(view_box[2]), float(view_box[3])
        except ValueError:
            width = height = None
    if not width or not height:
        return (1600.0, 900.0)
    return (width, height)


_svg_raster_cache: dict[tuple[Path, str | None], Path | None] = {}
_svg_raster_dir: tempfile.TemporaryDirectory | None = None


def _rasterize_svg(path: Path, heading: str | None = None) -> Path | None:
    """PNG rendering of an SVG at 2x, made with the headless Chromium the PDF export uses.

    PPTX has no portable SVG support, so figures are embedded as bitmaps. Returns
    None when no Chromium-based browser is installed; the caller skips the image.
    """
    global _svg_raster_dir
    key = (path, heading)
    if key in _svg_raster_cache:
        return _svg_raster_cache[key]
    from colloquium.export import _find_browser, _is_chromium_based

    out: Path | None = None
    browser = _find_browser()
    if browser and _is_chromium_based(browser):
        if _svg_raster_dir is None:
            _svg_raster_dir = tempfile.TemporaryDirectory(prefix="colloquium-svg-")
        width, height = _svg_size(path)
        scale = min(1.0, 1600.0 / width)
        w_px, h_px = max(1, round(width * scale)), max(1, round(height * scale))
        stem = f"{path.stem}-{len(_svg_raster_cache)}"
        source = path
        if heading:
            from colloquium.build import retitle_svg

            try:
                retitled = retitle_svg(path.read_text(encoding="utf-8"), html_module.escape(heading, quote=False))
            except OSError:
                retitled = None
            if retitled is not None:
                source = Path(_svg_raster_dir.name) / f"{stem}.svg"
                source.write_text(retitled, encoding="utf-8")
        wrapper = Path(_svg_raster_dir.name) / f"{stem}.html"
        wrapper.write_text(
            "<!doctype html><html><head><meta charset='utf-8'><style>"
            "html,body{margin:0;background:transparent}"
            f"img{{display:block;width:{w_px}px;height:{h_px}px}}"
            f"</style></head><body><img src=\"{html_module.escape(source.resolve().as_uri(), quote=True)}\"></body></html>",
            encoding="utf-8",
        )
        png = Path(_svg_raster_dir.name) / f"{stem}.png"
        cmd = [
            browser, "--headless=new", "--disable-gpu", "--hide-scrollbars",
            "--force-device-scale-factor=2", "--default-background-color=00000000",
            f"--screenshot={png}", f"--window-size={w_px},{h_px}", wrapper.as_uri(),
        ]
        try:
            subprocess.run(cmd, capture_output=True, timeout=120, check=False)
        except (OSError, subprocess.TimeoutExpired):
            pass
        if png.exists() and png.stat().st_size > 0:
            out = png
    _svg_raster_cache[key] = out
    return out


def _reset_svg_cache() -> None:
    global _svg_raster_dir
    _svg_raster_cache.clear()
    if _svg_raster_dir is not None:
        _svg_raster_dir.cleanup()
        _svg_raster_dir = None


def _image_aspect(path: Path) -> float:
    if path.suffix.lower() == ".svg":
        width, height = _svg_size(path)
        return width / height
    from PIL import Image

    with Image.open(path) as img:
        return img.width / img.height


def _layout_region(blocks, x, y, w, h, metrics: _Metrics, scale: float, gap: float):
    """Stack blocks top to bottom. Returns (placements, fits)."""
    min_img_h = min(1.2, h * 0.35)
    text_h = 0.0
    n_images = 0
    for block in blocks:
        if isinstance(block, TextBlock):
            text_h += metrics.text_height(block, w, scale)
        else:
            n_images += 1
            if block.legend:
                text_h += metrics.text_height(block.legend, w, scale) + gap * 0.4 + _CARD_PAD
    gaps = gap * max(0, len(blocks) - 1)
    remaining = h - text_h - gaps
    fits = remaining >= (min_img_h * n_images if n_images else 0)
    img_budget = max(remaining, min_img_h * n_images) / n_images if n_images else 0

    placements: list[_Placed] = []
    cursor = y
    for block in blocks:
        if isinstance(block, TextBlock):
            bh = metrics.text_height(block, w, scale)
            placements.append(_Placed("text", block, x, cursor, w, bh))
            cursor += bh + gap
        elif isinstance(block, DiagramBlock):
            diagram = block.diagram
            aspect = diagram.w / diagram.h
            ih = min(img_budget, w / aspect, diagram.h * block.max_scale)
            iw = ih * aspect
            placements.append(_Placed("diagram", block, x + (w - iw) / 2, cursor, iw, ih))
            cursor += ih + gap
        else:
            aspect = _image_aspect(block.path)
            ih = min(img_budget, w / aspect)
            iw = ih * aspect
            placements.append(_Placed("image", block, x + (w - iw) / 2, cursor, iw, ih))
            cursor += ih
            if block.legend:
                lh = metrics.text_height(block.legend, w, scale)
                cursor += gap * 0.4 + _CARD_PAD
                # Spans the column, which the figure is centered in, so a centered
                # caption sits under the figure and a right-aligned legend at its edge.
                right = block.legend.paras[0].align == "right"
                placements.append(_Placed("text", block.legend, x, cursor, (w + iw) / 2 if right else w, lh))
                cursor += lh
            cursor += gap
    return placements, fits


# ===== PPTX writing =====


def _set_run(run, spec: Run, size_pt: float, cfg: TemplateConfig, inherit_font: bool):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    run.text = spec.text
    font = run.font
    if not inherit_font:
        font.size = Pt(size_pt)
        if cfg.font:
            font.name = cfg.font
        color = spec.color or cfg.colors["text"]
        font.color.rgb = RGBColor.from_string(color)
    elif spec.color:
        font.color.rgb = RGBColor.from_string(spec.color)
    if spec.bold:
        font.bold = True
    if spec.italic:
        font.italic = True
    if spec.code:
        font.name = "Courier New"
    if spec.href:
        run.hyperlink.address = spec.href
        font.color.rgb = RGBColor.from_string(cfg.colors["link"])


def _fill_paragraph(p, runs: list[Run], size_pt: float, cfg: TemplateConfig, inherit_font: bool = False):
    for spec in runs:
        if spec.line_break:
            p.add_line_break()
            continue
        if spec.text:
            _set_run(p.add_run(), spec, size_pt, cfg, inherit_font)


def _set_bullet(p, para: Para, metrics: _Metrics, scale: float, cfg: TemplateConfig, space_before_pt: float):
    from lxml import etree
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    pPr = p._p.get_or_add_pPr()
    for child in list(pPr):
        pPr.remove(child)
    if space_before_pt:
        spc = etree.SubElement(pPr, qn("a:spcBef"))
        etree.SubElement(spc, qn("a:spcPts")).set("val", str(int(space_before_pt * 100)))
    if para.kind != "li":
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))
        return
    indent = Inches(metrics.indent_in(para, scale))
    hang = Inches(0.3 * scale * metrics.base_pt / 13.5)
    pPr.set("marL", str(int(indent)))
    pPr.set("indent", str(-int(hang)))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", cfg.colors["accent"])
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
    if para.ordered:
        font = etree.SubElement(pPr, qn("a:buFont"))
        font.set("typeface", cfg.font or "+mj-lt")
        num = etree.SubElement(pPr, qn("a:buAutoNum"))
        num.set("type", "arabicPeriod")
        if para.number != 1:
            num.set("startAt", str(para.number))
    else:
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•" if para.level == 0 else "–")


def _add_text_block(slide, block: TextBlock, x, y, w, h, metrics: _Metrics, scale: float, cfg: TemplateConfig):
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Inches

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(max(h, 0.2)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    prev = None
    for i, para in enumerate(block.paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        size = metrics.size(para, scale)
        _set_bullet(p, para, metrics, scale, cfg, metrics.space_before(para, prev, scale))
        p.alignment = {"right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}.get(para.align, PP_ALIGN.LEFT)
        runs = para.runs
        if para.kind == "h":
            runs = [Run(r.text, bold=True, italic=r.italic, code=r.code, href=r.href,
                        color=r.color, line_break=r.line_break) for r in runs]
        if para.kind == "small" and not any(r.color for r in runs):
            runs = [Run(r.text, r.bold, r.italic, r.code, r.href, cfg.colors["muted"], r.line_break)
                    for r in runs]
        _fill_paragraph(p, runs, size, cfg)
        prev = para
    return box


def _add_image(slide, placed: _Placed, cfg: TemplateConfig) -> bool:
    """Place a figure; False when an SVG could not be rasterized (no Chromium)."""
    from PIL import Image
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    if cfg.image_background:
        pad = _CARD_PAD
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(placed.x - pad), Inches(placed.y - pad),
            Inches(placed.w + 2 * pad), Inches(placed.h + 2 * pad),
        )
        card.adjustments[0] = 0.03
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor.from_string(cfg.image_background)
        card.line.fill.background()
        card.shadow.inherit = False
    path = placed.block.path
    if path.suffix.lower() == ".svg":
        png = _rasterize_svg(path, placed.block.heading)
        if png is None:
            return False
        source = str(png)
    else:
        with Image.open(path) as img:
            if (img.format or "").upper() in {"BMP", "GIF", "JPEG", "PNG", "TIFF"}:
                source = str(path)
            else:
                source = io.BytesIO()
                img.save(source, format="PNG")
                source.seek(0)
    slide.shapes.add_picture(source, Inches(placed.x), Inches(placed.y), Inches(placed.w), Inches(placed.h))
    return True


def _place_full_diagrams(entries, slide_w: float, cfg: TemplateConfig) -> None:
    """Draw slides whose body is a single diagram.

    Consecutive slides showing the same diagram (equal root tag and classes)
    are steps of one build-up: they share the largest scale that fits every
    step's body region and sit at the same spot, so nothing jumps between
    them, and each step after the first gets a Morph transition. The scale
    fits what the steps draw together, not the root box, so a diagram on a
    mostly empty canvas still fills the region, up to _DIAGRAM_MAX_ZOOM of
    its HTML size.
    """
    groups: list[list] = []
    for entry in entries:
        index, _, block = entry[:3]
        prev = groups[-1][-1] if groups else None
        if prev and prev[0] == index - 1 and prev[2].diagram.signature == block.diagram.signature:
            groups[-1].append(entry)
        else:
            groups.append([entry])
    for group in groups:
        top = max(e[4] for e in group)
        bottom = min(e[4] + e[6] for e in group)
        region_w = min(e[5] for e in group)
        bounds = [content_bounds(e[2].diagram) for e in group]
        bx = min(b[0] for b in bounds)
        by = min(b[1] for b in bounds)
        bw = max(b[0] + b[2] for b in bounds) - bx
        bh = max(b[1] + b[3] for b in bounds) - by
        scale = min(region_w / bw, (bottom - top) / bh, _DIAGRAM_MAX_ZOOM * slide_w / 1280)
        for step, (_, slide, block, margin_x, _, _, _) in enumerate(group):
            diagram = block.diagram
            x0 = margin_x + (region_w - bw * scale) / 2 - (bx - diagram.x) * scale
            y0 = top + max(0.0, (bottom - top - bh * scale) / 2) - (by - diagram.y) * scale
            add_diagram(slide, diagram, x0, y0, scale, cfg)
            if step:
                add_morph_transition(slide)


def _placeholder_font_pt(ph, default: float) -> float:
    """Best-effort font size for a placeholder, read from its layout/master XML."""
    from pptx.oxml.ns import qn

    candidates = [ph._element]
    try:
        layout_ph = ph._base_placeholder
        if layout_ph is not None:
            candidates.append(layout_ph._element)
            master_ph = layout_ph._base_placeholder
            if master_ph is not None:
                candidates.append(master_ph._element)
    except Exception:
        pass
    for el in candidates:
        for tag in ("a:defRPr", "a:rPr", "a:endParaRPr"):
            for node in el.iter(qn(tag)):
                sz = node.get("sz")
                if sz:
                    return int(sz) / 100
    return default


def _template_body_pt(layout) -> float | None:
    """Font size of the master's body placeholder, if the template sets one."""
    from pptx.enum.shapes import PP_PLACEHOLDER

    for ph in layout.slide_master.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.BODY:
            size = _placeholder_font_pt(ph, 0)
            return size or None
    return None


def _set_placeholder_runs(ph, runs: list[Run], cfg: TemplateConfig, all_caps: bool = False):
    tf = ph.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    _fill_paragraph(p, runs, 0, cfg, inherit_font=True)
    if all_caps:
        for r in p.runs:
            r.font._rPr.set("cap", "all")


def _widen_to_fit(ph, runs: list[Run], slide_w: float, char_em: float = 0.62):
    """Grow a short-text placeholder (speaker name, affiliation) rightward when
    the content is longer than the template's sample text, rather than letting
    it wrap into a cramped column."""
    from pptx.util import Emu, Inches

    size = _placeholder_font_pt(ph, 14)
    text = "".join("\n" if r.line_break else r.text for r in runs)
    longest = max((len(line) for line in text.split("\n")), default=0)
    needed = longest * char_em * size / 72 + 0.1
    left, top, width, height = ph.left, ph.top, ph.width, ph.height
    if left is None or width is None or needed <= Emu(width).inches:
        return
    max_width = slide_w - Emu(left).inches - 0.5
    _pin_geometry(ph, left, top, Emu(width), height)
    ph.width = Inches(max(Emu(width).inches, min(needed, max_width)))


def _shrink_to_fit(ph, runs: list[Run], bottom: float):
    """Shrink a placeholder's text (a multi-line subtitle on a closing slide,
    say) so its lines don't wrap, when a modest shrink manages that, and so it
    ends above ``bottom`` inches."""
    from pptx.util import Emu, Pt

    if ph.top is None or ph.width is None:
        return
    size = _placeholder_font_pt(ph, 18)
    text = "".join("\n" if r.line_break else r.text for r in runs)
    width_in = Emu(ph.width).inches
    room = bottom - Emu(ph.top).inches
    n_lines = text.count("\n") + 1

    def fits(s, wraps):
        lines = _estimate_lines(text, s, width_in, _BODY_CHAR_EM)
        return lines <= n_lines + wraps and lines * s * 1.2 / 72 <= room

    # Like headings: keep each source line on one line if a modest shrink does it.
    fitted = next((s for s in range(int(size), int(size * 0.72) - 1, -1) if fits(s, 0)), None)
    if fitted is None:
        fitted = next((s for s in range(int(size), int(size * 0.5) - 1, -1) if fits(s, n_lines)), int(size * 0.5))
    if fitted != size:
        for p in ph.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(fitted)


def _pin_geometry(ph, left, top, width, height):
    """Give a placeholder its own xfrm. Resizing an inheriting placeholder in
    python-pptx writes a fresh xfrm at 0,0 with zero size, so set all four."""
    ph.left, ph.top, ph.width, ph.height = left, top, width, height


def _plain_runs(markdown: str) -> list[Run]:
    md = create_base_md()
    tokens = md.parse(markdown.strip())
    runs: list[Run] = []
    for tok in tokens:
        if tok.type == "inline":
            if runs:
                runs.append(Run("", line_break=True))
            runs.extend(_inline_runs(tok.children))
    return runs


def _html_inline_runs(fragment: str) -> list[Run]:
    """Inline HTML/markdown fragment (e.g. title meta) -> runs."""
    text = html_module.unescape(_TAG_RE.sub("", fragment)).strip()
    return _plain_runs(text) if text else []


def _split_title_components(content: str) -> tuple[dict[str, list[Run]], str]:
    """Pull colloquium title-slide components out of a title slide's body."""
    found: dict[str, list[Run]] = {}
    for role, rx in _TITLE_COMPONENT_RES.items():
        m = rx.search(content)
        if m:
            found[role] = _html_inline_runs(m.group(1))
            content = content.replace(m.group(0), "")
    meta = _TITLE_META_RE.search(content)
    if meta:
        lines = [_html_inline_runs(p) for p in _HTML_P_RE.findall(meta.group(1))]
        lines = [line for line in lines if line]
        if lines:
            runs: list[Run] = []
            for line in lines:
                if runs:
                    runs.append(Run("", line_break=True))
                runs.extend(line)
            found["affiliation"] = runs
        content = content.replace(meta.group(0), "")
    content = _TITLE_RULE_RE.sub("", content)
    return found, content.strip()


def _cite_label_text(keys: list[str], bib_entries: dict, style: str, numbers: dict[str, int]) -> str:
    from colloquium.build import _format_citation_label

    labels = []
    for key in keys:
        if key in bib_entries:
            labels.append(_format_citation_label(bib_entries[key], key, style, numbers.get(key, 0)))
    return "; ".join(labels)


def _resolve_citations(text: str, bib_entries: dict, style: str, numbers: dict[str, int]) -> str:
    def _replace(m):
        keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
        label = _cite_label_text(keys, bib_entries, style, numbers)
        if not label:
            return m.group(0)
        return f"[{label}]" if style == "numeric" else f"({label})"

    return _CITATION_RE.sub(_replace, text)


def _fit_heading_size(title_lines: list[str], size: float, width_in: float, max_h: float | None = None) -> float:
    """Font size for a heading placeholder.

    Content headings (no ``max_h``) shrink, down to _HEADING_MIN_FRACTION of
    the template size, until each `##` line fits on one line, or failing that
    until the heading wraps at most once more. Full-bleed headings (title,
    section break) shrink until they fit ``max_h`` inches.
    """
    text = "\n".join(title_lines)

    def lines_at(s):
        return _estimate_lines(text, s, width_in, _HEADING_CHAR_EM)

    if max_h is not None:
        s = size
        while s > size * 0.4 and lines_at(s) * s * 1.15 / 72 > max_h:
            s -= 1
        return s
    steps = int((1 - _HEADING_MIN_FRACTION) / 0.02)
    candidates = [round(size * (1 - 0.02 * k) * 2) / 2 for k in range(steps + 1)]
    for limit in (len(title_lines), len(title_lines) + 1):
        for s in candidates:
            if lines_at(s) <= limit:
                return s
    return candidates[-1]


def _remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def _layout_spec_for(slide, cfg: TemplateConfig) -> tuple[str, LayoutSpec]:
    if slide.layout in cfg.layouts:
        return slide.layout, cfg.layouts[slide.layout]
    if slide.is_title_slide and "title" in cfg.layouts:
        return "title", cfg.layouts["title"]
    return "content", cfg.layouts["content"]


def _footer_text(value, n: int, total: int) -> str:
    text = str(value or "")
    if text == "auto":
        return f"{n} / {total}"
    return text.replace("{n}", str(n)).replace("{N}", str(total))


def export_pptx_with_template(deck, base_dir: Path, output_path: str, cfg: TemplateConfig) -> str:
    """Build ``deck`` onto the template in ``cfg`` and write ``output_path``."""
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    from colloquium.build import (
        _format_reference,
        _ordered_citation_keys,
        _parse_bib_file,
        expand_toc_markdown,
        slides_for_format,
    )

    prs = Presentation(str(cfg.template))
    template_slides = list(prs.slides)
    for name, spec in cfg.layouts.items():
        if not 1 <= spec.slide <= len(template_slides):
            raise ValueError(
                f"pptx.layouts.{name}: template has {len(template_slides)} slides, got slide {spec.slide}"
            )
    layout_of = {i + 1: s.slide_layout for i, s in enumerate(template_slides)}

    # Drop the template's sample slides; their media goes with them on save.
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)

    slide_w = Emu(prs.slide_width).inches
    slide_h = Emu(prs.slide_height).inches
    # Body text: the template's body size, kept within 1-1.2x the HTML's 24px.
    html_pt = 24 * slide_w * 72 / 1280
    body_pt = _template_body_pt(layout_of[cfg.layouts["content"].slide]) or html_pt
    metrics = _Metrics(base_pt=min(max(body_pt, html_pt), html_pt * 1.2))
    warnings: list[str] = []

    bib_entries = _parse_bib_file(deck.bibliography) if deck.bibliography else {}
    style = deck.citation_style
    slides = expand_toc_markdown(slides_for_format(deck.slides, "pptx"), deck.chapters)
    main = [s for s in slides if s.metadata.get("after") != "references"]
    post = [s for s in slides if s.metadata.get("after") == "references"]

    cited: list[str] = []
    if bib_entries:
        # Discover in rendered order so numeric citations follow first appearance.
        for s in main + post:
            keys = s.metadata.get("cite_left", []) + s.metadata.get("cite_right", [])
            for m in _CITATION_RE.finditer(s.content + " " + s.title):
                keys += [k.strip().lstrip("@") for k in m.group(1).split(";")]
            for key in keys:
                if key in bib_entries and key not in cited:
                    cited.append(key)
    numbers = {key: i + 1 for i, key in enumerate(cited)}
    ordered_refs = _ordered_citation_keys(cited, bib_entries, style, deck.citation_order, numbers) if cited else []

    from colloquium.slide import Slide

    ref_slides = []
    if ordered_refs:
        lines = []
        for key in ordered_refs:
            ref_html = _format_reference(bib_entries[key], key, style, numbers[key])
            url_match = re.search(r'href="([^"]+)"', ref_html)
            text = html_module.unescape(_TAG_RE.sub("", ref_html.replace("[link]", ""))).strip()
            lines.append(f"- {text} [link]({url_match.group(1)})" if url_match else f"- {text}")
        ref_slides.append(Slide(title="References", content="\n".join(lines)))

    ordered_slides = main + ref_slides + post
    # Like the HTML footer, {N} counts main + reference slides; appendix
    # slides after the references are numbered past it.
    total = len(main) + len(ref_slides)
    used_layouts = set()
    diagrams = snapshot_diagrams(deck, main + post, base_dir, warnings)
    # Slides whose whole body is one diagram, placed after the loop so that
    # consecutive steps of the same diagram share one scale and position.
    full_diagrams: list[tuple[int, object, DiagramBlock, float, float, float, float]] = []

    for index, src in enumerate(ordered_slides):
        layout_name, spec = _layout_spec_for(src, cfg)
        layout = layout_of[spec.slide]
        used_layouts.add(id(layout))
        slide = prs.slides.add_slide(layout)
        placeholders = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        roles = dict(spec.roles)
        roles.setdefault("heading", 0)
        filled: set[int] = set()

        content = src.content
        if bib_entries:
            content = _resolve_citations(content, bib_entries, style, numbers)

        # The parser joins consecutive leading `##` lines into one title with newlines;
        # each becomes its own line in the placeholder.
        title_lines = [line for line in src.title.split("\n") if line.strip()] if src.title else []
        while True:
            m = re.match(r"\s*##\s+(.*)(?:\n|$)", content)
            if not m or not title_lines:
                break
            title_lines.append(m.group(1).strip())
            content = content[m.end():]

        title_parts, content = _split_title_components(content) if src.is_title_slide else ({}, content)
        if src.metadata.get("eyebrow"):
            title_parts["eyebrow"] = _plain_runs(src.metadata["eyebrow"])
        if title_parts.get("note"):
            title_parts["subtitle"] = title_parts.pop("note")
        photo = None
        if src.is_title_slide and "photo" in roles:
            m = _MD_IMAGE_RE.search(content)
            if m:
                photo = _resolve_local(m.group(2), base_dir)
                content = content.replace(m.group(0), "").strip()
        if src.is_title_slide and content and "subtitle" in roles and "subtitle" not in title_parts:
            title_parts["subtitle"] = _plain_runs(_HTML_COMMENT_RE.sub("", content))
            content = ""

        heading_ph = placeholders.get(roles["heading"])
        heading_box = None
        if heading_ph is not None and heading_ph.left is not None:
            heading_box = (Emu(heading_ph.left).inches, Emu(heading_ph.top).inches)
        # figure-title: heading puts the heading into an SVG plot's own title, so the
        # plot carries it (as in HTML) and takes the heading's space.
        figure_owns_title = (
            src.metadata.get("figure_title") == "heading"
            and not src.is_title_slide
            and re.search(r"\.svg[\"')\s]", _HTML_COMMENT_RE.sub("", content)) is not None
        )
        if heading_ph is not None and title_lines and not figure_owns_title:
            runs: list[Run] = []
            for line in title_lines:
                if runs:
                    runs.append(Run("", line_break=True))
                runs.extend(_plain_runs(line) if line else [])
            _set_placeholder_runs(heading_ph, runs, cfg)
            filled.add(roles["heading"])
        elif title_lines and heading_ph is None:
            warnings.append(f"slide {index + 1}: layout '{layout_name}' has no heading placeholder idx {roles['heading']}")

        for role in ("eyebrow", "subtitle", "name", "affiliation"):
            runs = title_parts.pop(role, None)
            if not runs:
                continue
            ph = placeholders.get(roles.get(role, -1))
            if ph is None:
                # No slot for it on this layout: keep it as body text instead.
                content = "".join(r.text for r in runs) + ("\n\n" + content if content else "")
                continue
            _set_placeholder_runs(ph, runs, cfg, all_caps=(role == "eyebrow"))
            _widen_to_fit(ph, runs, slide_w)
            if role == "subtitle":
                _shrink_to_fit(ph, runs, slide_h - 0.35)
            filled.add(roles[role])
        photo_ph = placeholders.get(roles.get("photo", -1))
        if photo is None and photo_ph is not None and photo_ph.left is not None:
            # No headshot: the name and affiliation take the photo's spot instead of
            # sitting indented beside an empty one.
            for role in ("name", "affiliation"):
                ph = placeholders.get(roles.get(role, -1))
                if roles.get(role) in filled and ph.left is not None and ph.left > photo_ph.left:
                    _pin_geometry(ph, photo_ph.left, ph.top, ph.width, ph.height)
        if photo is not None and hasattr(placeholders.get(roles.get("photo", -1)), "insert_picture"):
            # insert_picture swaps the placeholder element for a filled picture.
            placeholders[roles["photo"]].insert_picture(str(photo))
            filled.add(roles["photo"])
            placeholders = {p.placeholder_format.idx: p for p in slide.placeholders}

        for idx, ph in placeholders.items():
            if idx not in filled:
                _remove_shape(ph)

        # ----- Body content region -----
        margin_x = heading_box[0] if heading_box else 0.6
        region_w = slide_w - 2 * margin_x
        top = heading_box[1] if heading_box and figure_owns_title else 0.6
        content = _HTML_COMMENT_RE.sub("", content).strip()
        is_full_bleed = src.is_title_slide or layout_name == "section-break"
        if heading_ph is not None and roles["heading"] in filled:
            heading_top = heading_ph.top
            if not is_full_bleed:
                # Content headings span the slide.
                _pin_geometry(heading_ph, heading_ph.left, heading_top, Inches(region_w), heading_ph.height)
            template_size = _placeholder_font_pt(heading_ph, 30)
            width_in = Emu(heading_ph.width).inches
            max_h = None
            if is_full_bleed:
                max_h = Emu(heading_ph.height).inches
                if layout_name == "section-break":
                    # Nothing sits under a section title: it may use the slide down to the margin.
                    max_h = max(max_h, slide_h - Emu(heading_top).inches - 0.4)
            size = _fit_heading_size(title_lines, template_size, width_in, max_h)
            if size != template_size:
                for p in heading_ph.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(size)
            text = "\n".join(title_lines)
            lines = _estimate_lines(text, size, width_in, _HEADING_CHAR_EM)
            title_h = lines * size * 1.15 / 72 + 0.08
            if not is_full_bleed:
                heading_ph.height = Inches(title_h)
            top = Emu(heading_top).inches + title_h + 0.22

        has_cite = bool(src.metadata.get("cite_left") or src.metadata.get("cite_right"))
        bottom_reserved = 0.5 + (0.22 if has_cite else 0)
        region_h = slide_h - top - bottom_reserved

        if content:
            columns = [c.strip() for c in _COLUMN_SPLIT_RE.split(content)]
            fracs = None
            cols_class = next((c for c in src.classes if c.startswith("cols-")), None)
            if cols_class and len(columns) > 1:
                parts = [float(v) for v in cols_class[5:].split("-") if v]
                if len(parts) == len(columns):
                    fracs = [v / sum(parts) for v in parts]
            fracs = fracs or [1 / len(columns)] * len(columns)
            col_gap = 0.4 if len(columns) > 1 else 0
            usable = region_w - col_gap * (len(columns) - 1)
            diagram_numbers = iter(range(len(content)))

            def next_diagram(src=src):
                diagram = diagrams.get((id(src), next(diagram_numbers)))
                return DiagramBlock(diagram, max_scale=slide_w / 1280) if diagram else None

            col_blocks = [_parse_blocks(c, base_dir, warnings, next_diagram) if c else [] for c in columns]
            all_blocks = [b for blocks in col_blocks for b in blocks]
            if len(all_blocks) == 1 and isinstance(all_blocks[0], DiagramBlock):
                full_diagrams.append((index, slide, all_blocks[0], margin_x, top, region_w, region_h))
                col_blocks = []
            if src.metadata.get("figure_title") == "heading" and title_lines:
                first_image = next((b for blocks in col_blocks for b in blocks if isinstance(b, ImageBlock)), None)
                if first_image is not None and first_image.path.suffix.lower() == ".svg":
                    first_image.heading = " ".join(title_lines)

            scale = 1.0
            while True:
                gap = 0.16 * scale
                all_fit = True
                plan = []
                x = margin_x
                for blocks, frac in zip(col_blocks, fracs):
                    w = usable * frac
                    placed, fits = _layout_region(blocks, x, top, w, region_h, metrics, scale, gap)
                    used_h = max((p.y + p.h for p in placed), default=top) - top
                    all_fit &= fits and used_h <= region_h + 0.01
                    plan.extend(placed)
                    x += w + col_gap
                if all_fit or scale <= _MIN_SCALE:
                    if not all_fit:
                        warnings.append(f"slide {index + 1} ('{src.title}'): content may overflow; consider splitting it")
                    break
                scale = max(_MIN_SCALE, scale * 0.94)

            for placed in plan:
                if placed.kind == "diagram":
                    diagram = placed.block.diagram
                    add_diagram(slide, diagram, placed.x, placed.y, placed.w / diagram.w, cfg)
                elif placed.kind == "image":
                    if not _add_image(slide, placed, cfg):
                        warnings.append(
                            f"slide {index + 1}: {placed.block.path.name} needs a Chromium-based browser to rasterize SVG; skipped"
                        )
                else:
                    _add_text_block(slide, placed.block, placed.x, placed.y, placed.w, placed.h, metrics, scale, cfg)

        # ----- Citations and footer (content slides only) -----
        cite_keys = src.metadata.get("cite_left", []) + src.metadata.get("cite_right", [])
        if cite_keys and bib_entries:
            label = _cite_label_text(cite_keys, bib_entries, style, numbers)
            if label:
                block = TextBlock([Para([Run(label)], kind="small")])
                _add_text_block(slide, block, margin_x, slide_h - 0.68, region_w, 0.2, metrics, 1.0, cfg)
        if deck.footer and not is_full_bleed:
            for zone, align in (("left", "left"), ("center", "center"), ("right", "right")):
                value = deck.footer.get(zone)
                if not value or re.search(r"\.(png|jpe?g|gif|svg|webp)$", str(value), re.IGNORECASE):
                    continue
                text = _footer_text(value, index + 1, total)
                para = Para([Run(text)], kind="small", align="right" if align == "right" else "left")
                box = _add_text_block(slide, TextBlock([para]), margin_x, slide_h - 0.4, region_w, 0.2, metrics, 1.0, cfg)
                if align == "center":
                    from pptx.enum.text import PP_ALIGN

                    box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        if src.speaker_notes:
            slide.notes_slide.notes_text_frame.text = src.speaker_notes

    _place_full_diagrams(full_diagrams, slide_w, cfg)

    # Prune layouts nothing uses so the template's sample backgrounds don't ship.
    for master in prs.slide_masters:
        for layout in list(master.slide_layouts):
            if id(layout) not in used_layouts and not layout.used_by_slides:
                master.slide_layouts.remove(layout)

    prs.core_properties.title = deck.title
    if deck.author:
        prs.core_properties.author = deck.author
    from colloquium.export import _register_notes_master

    _register_notes_master(prs)  # Keynote rejects decks whose notes master is unlisted
    prs.save(output_path)
    _reset_svg_cache()

    for warning in dict.fromkeys(warnings):
        print(f"Warning: {warning}", file=sys.stderr)
    return output_path
