"""HTML diagrams as native PPTX shapes, laid out by the browser.

A slide's raw HTML block that contains ``<div>`` or ``<svg>`` markup (boxes,
labels, connector lines) is a diagram: flattening it to text loses the
picture. Instead, the deck is built to HTML once and opened in headless
Chromium, which settles every animation and reports where each box, text
run, SVG line, and image ends up on the 1280x720 canvas. Those become
editable rectangles, text boxes, lines, and pictures on the PPTX slide.

Shapes are named from their place in the HTML (``!!`` prefix), so the same
element keeps the same name on consecutive slides. Runs of slides that show
the same diagram at different steps share one scale and position and get
PowerPoint's Morph transition, which animates matching shapes between them,
much like the deck's CSS animations do in the browser.
"""

from __future__ import annotations

import base64
import copy
import hashlib
import json
import math
import re
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import unquote, urlsplit

from colloquium.md import create_base_md

DIAGRAM_ATTR = "data-colloquium-diagram"
_SLIDE_W_PX = 1280

_COMMENT_RE = re.compile(r"<!--.*?-->", re.DOTALL)
_STRUCTURE_RE = re.compile(r"<(div|svg)\b", re.IGNORECASE)
_FIRST_TAG_RE = re.compile(r"<([a-zA-Z][\w:-]*)")
_LAYOUT_JSON_RE = re.compile(
    r'<script type="application/json" id="colloquium-diagram-layout">(.*?)</script>', re.DOTALL
)

# Theme variables whose colors (and their tints and shades) follow the template palette.
_ROLE_VARS = {
    "text": ("--colloquium-text", "--colloquium-heading"),
    "muted": ("--colloquium-muted",),
    "accent": ("--colloquium-accent",),
    "link": ("--colloquium-link",),
}
_MORPH_MS = 700


def is_diagram_html(html: str) -> bool:
    """True for an HTML block with layout structure: an ``<svg>`` or nested ``<div>``s.

    A single ``<div>`` of inline spans (a plot legend, a caption) stays text,
    which the template exporter already renders, swatches included.
    """
    tags = [t.lower() for t in _STRUCTURE_RE.findall(_COMMENT_RE.sub("", html))]
    return "svg" in tags or tags.count("div") >= 2


def _mark_diagrams(content: str, slide_key: int) -> tuple[str, int]:
    """Tag the root element of each diagram HTML block in ``content``.

    Returns the new content and how many diagrams it holds. Diagrams are
    numbered in source order, the same order the PPTX exporter meets them.
    """
    lines = content.split("\n")
    count = 0
    for tok in create_base_md().parse(content):
        if tok.type != "html_block" or not tok.map or not is_diagram_html(tok.content):
            continue
        start, end = tok.map
        block = "\n".join(lines[start:end])
        stripped = _COMMENT_RE.sub(lambda m: " " * len(m.group(0)), block)
        m = _FIRST_TAG_RE.search(stripped)
        if m is None:
            continue
        at = m.end()
        block = f'{block[:at]} {DIAGRAM_ATTR}="{slide_key}:{count}"{block[at:]}'
        lines[start:end] = block.split("\n")
        count += 1
    return "\n".join(lines), count


@dataclass
class Diagram:
    """A diagram's root box and its settled contents, in CSS px of the slide canvas."""

    signature: str  # root tag and classes; equal on every step of one diagram
    x: float
    y: float
    w: float
    h: float
    items: list[dict] = field(default_factory=list)
    theme: dict[str, list] = field(default_factory=dict)


def snapshot_diagrams(deck, slides: list, base_dir: Path, warnings: list[str]) -> dict[tuple[int, int], Diagram]:
    """Lay out every diagram in ``slides`` in headless Chromium.

    Returns ``{(id(slide), n): Diagram}`` for the n-th diagram of each slide.
    Empty when no slide has a diagram or no Chromium-based browser exists.
    """
    from dataclasses import replace

    from colloquium.build import build_deck
    from colloquium.export import _find_browser, _is_chromium_based, _pdf_time_budget_ms, _serve_html_for_export

    keys: dict[int, int] = {}
    marked = []
    for slide in slides:
        # Title slides' HTML is title components, which fill template placeholders.
        content, count = (slide.content, 0) if slide.is_title_slide else _mark_diagrams(slide.content, len(keys))
        if count:
            keys[id(slide)] = len(keys)
        metadata = {k: v for k, v in slide.metadata.items() if k != "only"}
        marked.append(replace(slide, content=content, metadata=metadata))
    if not keys:
        return {}

    browser = _find_browser()
    if browser is None or not _is_chromium_based(browser):
        warnings.append("HTML diagrams need a Chromium-based browser to lay out; exported as text instead")
        return {}

    snapshot_deck = copy.copy(deck)
    snapshot_deck.slides = marked
    html = build_deck(snapshot_deck, output_stem="colloquium-pptx-diagrams")
    script = f"<script>{_MEASURE_JS}</script>"
    html = html.replace("</body>", f"{script}</body>") if "</body>" in html else html + script

    # Next to the source so relative assets and custom CSS resolve, removed afterwards.
    with tempfile.NamedTemporaryFile(
        "w", suffix=".html", prefix=".colloquium-pptx-", dir=base_dir, delete=False, encoding="utf-8"
    ) as tmp:
        tmp.write(html)
        tmp_path = Path(tmp.name)
    try:
        with _serve_html_for_export(str(tmp_path)) as url:
            cmd = [
                browser,
                "--headless=new",
                "--disable-gpu",
                "--hide-scrollbars",
                "--window-size=1280,720",
                f"--virtual-time-budget={_pdf_time_budget_ms()}",
                "--dump-dom",
                f"{url}?capture",
            ]
            try:
                result = subprocess.run(cmd, capture_output=True, text=True, timeout=120, check=False)
            except (OSError, subprocess.TimeoutExpired):
                result = None
            origin = "{0.scheme}://{0.netloc}".format(urlsplit(url))
    finally:
        tmp_path.unlink(missing_ok=True)

    m = _LAYOUT_JSON_RE.search(result.stdout) if result is not None else None
    if m is None:
        warnings.append("headless browser did not lay out the HTML diagrams; exported as text instead")
        return {}
    data = json.loads(m.group(1))
    if data.get("error"):
        warnings.append(f"laying out HTML diagrams failed ({data['error']}); exported as text instead")
        return {}
    for tag in data.get("unsupported", []):
        warnings.append(f"diagram SVG <{tag}> elements are not converted to shapes; skipped")

    by_key = {v: k for k, v in keys.items()}
    out: dict[tuple[int, int], Diagram] = {}
    for key, entry in data.get("diagrams", {}).items():
        slide_key, n = (int(part) for part in key.split(":"))
        for item in entry["items"]:
            if item["kind"] == "image":
                item["path"] = _resolve_image(item.pop("src"), origin, base_dir, warnings)
        out[(by_key[slide_key], n)] = Diagram(
            signature=entry["signature"],
            x=entry["x"], y=entry["y"], w=entry["w"], h=entry["h"],
            items=entry["items"],
            theme=entry.get("theme", {}),
        )
    return out


def _resolve_image(src: str, origin: str, base_dir: Path, warnings: list[str]):
    if src.startswith("data:"):
        header, _, payload = src.partition(",")
        return base64.b64decode(payload) if header.endswith(";base64") else None
    if src.startswith(origin + "/"):
        path = base_dir / unquote(urlsplit(src).path).lstrip("/")
        if path.is_file():
            return path
    warnings.append(f"diagram image not found or remote, skipped: {src}")
    return None


# ===== Colors =====


def _palette(theme: dict[str, list], colors: dict[str, str]) -> list[tuple[tuple, tuple]]:
    pairs = []
    for role, names in _ROLE_VARS.items():
        target = colors.get(role)
        if not target:
            continue
        dst = tuple(int(target[i:i + 2], 16) for i in (0, 2, 4))
        for name in names:
            src = theme.get(name)
            if src:
                pairs.append((tuple(src[:3]), dst))
    return pairs


def map_color(rgb, pairs) -> tuple[int, int, int]:
    """Move a theme color, or a tint or shade of one, onto the template palette.

    ``rgb`` is matched against each theme color mixed with white or black; the
    closest fit within a small tolerance is redone with the template color at
    the same mix. Colors unrelated to the theme are kept.
    """
    rgb = tuple(rgb[:3])
    best = None
    for src, dst in pairs:
        for base in ((255, 255, 255), (0, 0, 0)):
            d = [s - b for s, b in zip(src, base)]
            norm = sum(v * v for v in d)
            if norm < 1:
                continue
            t = sum((c - b) * v for c, b, v in zip(rgb, base, d)) / norm
            if not 0.05 <= t <= 1.1:
                continue
            resid = math.dist(rgb, [b + t * v for b, v in zip(base, d)])
            if resid <= 16 and (best is None or resid < best[0]):
                best = (resid, [b + t * (o - b) for b, o in zip(base, dst)])
    if best is None:
        return tuple(int(v) for v in rgb[:3])
    return tuple(max(0, min(255, round(v))) for v in best[1])


# ===== Placement =====


@dataclass
class _Frame:
    """Maps diagram CSS px to slide inches: ``x0 + (x - root.x) * s``."""

    x0: float
    y0: float
    s: float
    root_x: float
    root_y: float

    def x(self, v):
        return self.x0 + (v - self.root_x) * self.s

    def y(self, v):
        return self.y0 + (v - self.root_y) * self.s

    def pt(self, px):
        return px * self.s * 72


def fit_scale(diagram: Diagram, region_w: float, region_h: float, slide_w: float) -> float:
    """Inches per CSS px: as large as the region allows, never above the HTML's own size."""
    return min(region_w / diagram.w, region_h / diagram.h, slide_w / _SLIDE_W_PX)


def shape_name(key: str) -> str:
    """Stable ``!!`` name from an element's path in the diagram, for Morph matching."""
    leaf = re.sub(r"[^a-z0-9-]+", "-", key.rsplit(">", 1)[-1].split("#")[0].lower()).strip("-")
    return f"!!{leaf[:24]}-{hashlib.sha1(key.encode()).hexdigest()[:6]}"


def add_diagram(slide, diagram: Diagram, x0: float, y0: float, s: float, cfg) -> None:
    """Draw ``diagram`` on ``slide`` with its root's top-left at (x0, y0) inches."""
    frame = _Frame(x0, y0, s, diagram.x, diagram.y)
    pairs = _palette(diagram.theme, cfg.colors)
    for item in sorted(diagram.items, key=lambda it: (it.get("z", 0), it["order"])):
        kind = item["kind"]
        if kind == "box":
            _add_box(slide, item, frame, pairs, cfg)
        elif kind == "text":
            _add_text(slide, item, frame, pairs, cfg)
        elif kind == "line":
            _add_line(slide, item, frame, pairs)
        elif kind == "image" and item.get("path") is not None:
            _add_picture(slide, item, frame)


def _rgb(color, pairs):
    from pptx.dml.color import RGBColor

    return RGBColor(*map_color(color, pairs))


def _set_alpha(fill_or_line_elm, alpha: float):
    """Add ``<a:alpha>`` to the solid fill color under ``fill_or_line_elm``."""
    from lxml import etree
    from pptx.oxml.ns import qn

    if alpha >= 0.999 or fill_or_line_elm is None:
        return
    clr = fill_or_line_elm.find(".//" + qn("a:srgbClr"))
    if clr is not None:
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(round(alpha * 100000))))


def _add_box(slide, item, frame: _Frame, pairs, cfg):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    w, h = item["w"], item["h"]
    radius = item.get("radius", 0)
    if radius >= min(w, h) / 2 - 0.5 and abs(w - h) < 1:
        kind = MSO_SHAPE.OVAL
    elif radius > 0.5:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE
    else:
        kind = MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        kind, Inches(frame.x(item["x"])), Inches(frame.y(item["y"])),
        Inches(max(w * frame.s, 0.01)), Inches(max(h * frame.s, 0.01)),
    )
    shape.name = shape_name(item["key"])
    shape.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = min(0.5, radius / min(w, h))
    opacity = item.get("opacity", 1)
    _set_fill(shape, item.get("fill"), opacity, pairs)
    if item.get("stroke"):
        shape.line.color.rgb = _rgb(item["stroke"], pairs)
        shape.line.width = Pt(frame.pt(item["stroke_width"]))
        _set_alpha(shape.line._get_or_add_ln(), item["stroke"][3] * opacity)
    else:
        shape.line.fill.background()

    if item.get("runs"):
        tf = shape.text_frame
        pad = item.get("padding", [0, 0, 0, 0])
        tf.margin_top, tf.margin_right, tf.margin_bottom, tf.margin_left = (Inches(v * frame.s) for v in pad)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}.get(item.get("valign"), MSO_ANCHOR.MIDDLE)
        _fill_text(tf, item, frame, pairs, cfg)


def _set_fill(shape, fill, opacity: float, pairs):
    from pptx.oxml.ns import qn

    if not fill:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill, pairs)
    _set_alpha(shape.fill._xPr.find(qn("a:solidFill")), fill[3] * opacity)


def _add_text(slide, item, frame: _Frame, pairs, cfg):
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.util import Inches

    box = slide.shapes.add_textbox(
        Inches(frame.x(item["x"])), Inches(frame.y(item["y"])),
        Inches(max(item["w"] * frame.s, 0.05)), Inches(max(item["h"] * frame.s, 0.05)),
    )
    box.name = shape_name(item["key"])
    tf = box.text_frame
    tf.auto_size = MSO_AUTO_SIZE.NONE
    # Centered on the text's measured middle, so a font that wraps differently
    # grows the text evenly up and down instead of off one edge.
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    _fill_text(tf, item, frame, pairs, cfg)


def _fill_text(tf, item, frame: _Frame, pairs, cfg):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    tf.word_wrap = item.get("wrap", True)
    p = tf.paragraphs[0]
    p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(item.get("align"), PP_ALIGN.LEFT)
    opacity = item.get("opacity", 1)
    for spec in item["runs"]:
        if spec.get("br"):
            p.add_line_break()
            continue
        run = p.add_run()
        run.text = spec["text"]
        font = run.font
        font.size = Pt(round(frame.pt(spec["size"]) * 2) / 2)
        if cfg.font:
            font.name = cfg.font
        font.bold = spec.get("bold", False)
        font.italic = spec.get("italic", False)
        if spec.get("color"):
            font.color.rgb = _rgb(spec["color"], pairs)
            _set_alpha(run._r.get_or_add_rPr(), spec["color"][3] * opacity)
        if spec.get("href"):
            run.hyperlink.address = spec["href"]


def _add_line(slide, item, frame: _Frame, pairs):
    from lxml import etree
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    points = item["points"]
    if len(points) == 2 and not item.get("closed"):
        (x1, y1), (x2, y2) = points
        shape = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(frame.x(x1)), Inches(frame.y(y1)), Inches(frame.x(x2)), Inches(frame.y(y2)),
        )
    else:
        builder = slide.shapes.build_freeform(Inches(frame.x(points[0][0])), Inches(frame.y(points[0][1])), scale=1.0)
        builder.add_line_segments(
            [(Inches(frame.x(x)), Inches(frame.y(y))) for x, y in points[1:]], close=bool(item.get("closed"))
        )
        shape = builder.convert_to_shape()
        _set_fill(shape, item.get("fill"), item.get("opacity", 1), pairs)
        shape.shadow.inherit = False
    shape.name = shape_name(item["key"])
    line = shape.line
    width = item.get("stroke_width", 1)
    if item.get("stroke"):
        line.color.rgb = _rgb(item["stroke"], pairs)
        line.width = Pt(frame.pt(width))
        _set_alpha(line._get_or_add_ln(), item["stroke"][3] * item.get("opacity", 1))
    else:
        line.fill.background()
    dash = item.get("dash")
    if dash and width > 0:
        ratio = dash[0] / width
        line.dash_style = (
            MSO_LINE_DASH_STYLE.ROUND_DOT if ratio <= 1.5
            else MSO_LINE_DASH_STYLE.DASH if ratio <= 4.5
            else MSO_LINE_DASH_STYLE.LONG_DASH
        )
    ln = line._get_or_add_ln()
    for end, tag in (("arrow_start", "a:headEnd"), ("arrow_end", "a:tailEnd")):
        if item.get(end):
            etree.SubElement(ln, qn(tag)).set("type", "triangle")


def _add_picture(slide, item, frame: _Frame):
    import io

    from pptx.util import Inches

    source = item["path"]
    if isinstance(source, bytes):
        source = io.BytesIO(source)
    x, y, w, h = item["x"], item["y"], item["w"], item["h"]
    nw, nh = item.get("natural_w") or 0, item.get("natural_h") or 0
    if item.get("fit") == "contain" and nw and nh:
        k = min(w / nw, h / nh)
        x, y, w, h = x + (w - nw * k) / 2, y + (h - nh * k) / 2, nw * k, nh * k
    pic = slide.shapes.add_picture(
        source if not isinstance(source, Path) else str(source),
        Inches(frame.x(x)), Inches(frame.y(y)), Inches(w * frame.s), Inches(h * frame.s),
    )
    pic.name = shape_name(item["key"])


def add_morph_transition(slide, duration_ms: int = _MORPH_MS) -> None:
    """Give ``slide`` PowerPoint's Morph transition, with a fade where Morph is unknown."""
    from lxml import etree

    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    xml = (
        f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        f'<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
        f'<p:transition xmlns:p="{p}" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
        f'spd="slow" p14:dur="{duration_ms}"><p159:morph option="byObject"/></p:transition>'
        f"</mc:Choice>"
        f'<mc:Fallback><p:transition xmlns:p="{p}" spd="med"><p:fade/></p:transition></mc:Fallback>'
        f"</mc:AlternateContent>"
    )
    sld = slide._element
    mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
    for old in sld.findall(f"{{{p}}}transition") + [
        alt for alt in sld.findall(mc) if alt.find(f".//{{{p}}}transition") is not None
    ]:
        sld.remove(old)
    # Schema order: cSld, clrMapOvr, transition, timing, extLst.
    anchor = sld.find(f"{{{p}}}clrMapOvr")
    if anchor is None:
        anchor = sld.find(f"{{{p}}}cSld")
    anchor.addnext(etree.fromstring(xml))


# Runs in the snapshot page: settles each marked diagram's slide and records its
# boxes, text, SVG lines, and images relative to the slide, in CSS px.
_MEASURE_JS = r"""
(() => {
  const ATTR = 'data-colloquium-diagram';
  const THEME_VARS = ['--colloquium-text', '--colloquium-heading', '--colloquium-muted',
                      '--colloquium-accent', '--colloquium-link', '--colloquium-border', '--colloquium-bg'];
  const SVG_SHAPES = new Set(['line', 'polyline', 'polygon', 'rect', 'circle', 'ellipse']);
  const SVG_SKIP = new Set(['svg', 'g', 'defs', 'title', 'desc', 'marker', 'clippath', 'mask', 'pattern',
                            'symbol', 'style', 'lineargradient', 'radialgradient', 'stop', 'use']);
  const px = (v) => parseFloat(v) || 0;
  const canvas = document.createElement('canvas');
  canvas.width = canvas.height = 1;
  const ctx = canvas.getContext('2d', {willReadFrequently: true});
  const colors = new Map();
  // Any CSS color syntax -> [r, g, b, a], or null when invisible.
  function color(s) {
    if (!s || s === 'none' || s === 'transparent' || s.startsWith('url(')) return null;
    if (colors.has(s)) return colors.get(s);
    ctx.clearRect(0, 0, 1, 1);
    ctx.fillStyle = 'rgba(0, 0, 0, 0)';
    ctx.fillStyle = s;
    ctx.fillRect(0, 0, 1, 1);
    const d = ctx.getImageData(0, 0, 1, 1).data;
    const out = d[3] ? [d[0], d[1], d[2], Math.round(d[3] / 255 * 1000) / 1000] : null;
    colors.set(s, out);
    return out;
  }
  function opacity(el, stop) {
    let o = 1;
    for (let n = el; n && n !== stop.parentElement; n = n.parentElement) o *= px(getComputedStyle(n).opacity);
    return o;
  }
  function zOf(el, root) {
    for (let n = el; n && n !== root.parentElement; n = n.parentElement) {
      const cs = getComputedStyle(n);
      if (cs.position !== 'static' && cs.zIndex !== 'auto') return parseInt(cs.zIndex, 10) || 0;
    }
    return 0;
  }
  function signature(el) {
    const tag = el.tagName.toLowerCase();
    const cls = [...el.classList].filter((c) => c !== 'active' && c !== 'visible').sort();
    return cls.length ? tag + '.' + cls.join('.') : tag;
  }
  // Path from the root, with an index among same-signature siblings.
  function keyOf(el, root) {
    const parts = [];
    for (let n = el; n && n !== root; n = n.parentElement) {
      const sig = signature(n);
      const twins = [...n.parentElement.children].filter((c) => signature(c) === sig);
      parts.unshift(twins.length > 1 ? sig + ':' + twins.indexOf(n) : sig);
    }
    return parts.join('>') || 'root';
  }
  function transform(text, mode) {
    if (mode === 'uppercase') return text.toUpperCase();
    if (mode === 'lowercase') return text.toLowerCase();
    if (mode === 'capitalize') return text.replace(/\b\w/g, (c) => c.toUpperCase());
    return text;
  }
  function runsOf(el, onlyDirect) {
    const runs = [];
    const walk = (node) => {
      for (const n of node.childNodes) {
        if (n.nodeType === 3) {
          const cs = getComputedStyle(n.parentElement);
          let t = n.data;
          if (!/^pre/.test(cs.whiteSpace)) t = t.replace(/[ \t\n\r]+/g, ' ');
          t = transform(t, cs.textTransform);
          if (!t) continue;
          const prev = runs[runs.length - 1];
          if (t.startsWith(' ') && (!prev || prev.br || prev.text.endsWith(' '))) t = t.slice(1);
          if (!t) continue;
          const a = n.parentElement.closest('a');
          runs.push({text: t, size: px(cs.fontSize), bold: (parseInt(cs.fontWeight, 10) || 400) >= 600,
                     italic: cs.fontStyle === 'italic', color: color(cs.color), href: a ? a.href : null});
        } else if (n.nodeType === 1 && !onlyDirect) {
          if (n.tagName === 'BR') runs.push({br: true});
          else if (getComputedStyle(n).display !== 'none') walk(n);
        }
      }
    };
    walk(el);
    while (runs.length && runs[runs.length - 1].br) runs.pop();
    const last = runs[runs.length - 1];
    if (last) last.text = last.text.replace(/ +$/, '');
    return runs.filter((r) => r.br || r.text);
  }
  function measure(root, unsupported) {
    const slide = root.closest('.slide');
    document.querySelectorAll('.slide.active').forEach((s) => s.classList.remove('active'));
    slide.classList.add('active');
    void slide.offsetWidth;
    // Settle every animation and transition at its end state, as the audience sees it.
    for (const a of document.getAnimations()) { try { a.finish(); } catch (e) { a.cancel(); } }
    const sr = slide.getBoundingClientRect();
    const k = sr.width / slide.offsetWidth || 1;
    const rel = (r) => ({x: (r.left - sr.left) / k, y: (r.top - sr.top) / k, w: r.width / k, h: r.height / k});
    const point = (m, x, y) => {
      const p = new DOMPoint(x, y).matrixTransform(m);
      return [(p.x - sr.left) / k, (p.y - sr.top) / k];
    };
    const items = [];
    // Elements whose text an ancestor's text box already carries.
    const consumed = new Set();
    let order = 0;
    for (const el of [root, ...root.querySelectorAll('*')]) {
      order++;
      const tag = el.tagName.toLowerCase();
      const inSvg = el.ownerSVGElement != null;
      if (consumed.has(el) && tag !== 'img') continue;
      if (el.getClientRects().length === 0) continue;
      const cs = getComputedStyle(el);
      if (cs.visibility === 'hidden') continue;
      const op = opacity(el, slide);
      if (op < 0.02) continue;
      const base = {key: keyOf(el, root), order, z: zOf(el, root), opacity: Math.round(op * 1000) / 1000};
      if (inSvg) {
        if (el.closest('defs, marker, clipPath, mask, pattern, symbol')) continue;
        if (!SVG_SHAPES.has(tag)) {
          if (!SVG_SKIP.has(tag)) unsupported.add(tag);
          continue;
        }
        const m = el.getScreenCTM();
        if (!m) continue;
        const scaleOf = Math.sqrt(Math.abs(m.a * m.d - m.b * m.c));
        const nonScaling = cs.vectorEffect === 'non-scaling-stroke';
        const stroke = color(cs.stroke);
        const sw = px(cs.strokeWidth) * (nonScaling ? 1 : scaleOf) / k;
        const dash = cs.strokeDasharray && cs.strokeDasharray !== 'none'
          ? cs.strokeDasharray.split(/[ ,]+/).map(px).map((v) => v * (nonScaling ? 1 : scaleOf) / k) : null;
        const fill = color(cs.fill);
        const g = (name) => el[name].baseVal.value;
        if (tag === 'line' || tag === 'polyline' || tag === 'polygon') {
          const pts = tag === 'line'
            ? [point(m, g('x1'), g('y1')), point(m, g('x2'), g('y2'))]
            : [...el.points].map((p) => point(m, p.x, p.y));
          if (pts.length < 2 || (!stroke && !fill)) continue;
          items.push({...base, kind: 'line', points: pts, closed: tag === 'polygon',
                      stroke, stroke_width: sw, dash, fill: tag === 'line' ? null : fill,
                      arrow_start: cs.markerStart !== 'none', arrow_end: cs.markerEnd !== 'none'});
        } else {
          let x0, y0, x1, y1, radius = 0;
          if (tag === 'rect') {
            [x0, y0] = point(m, g('x'), g('y'));
            [x1, y1] = point(m, g('x') + g('width'), g('y') + g('height'));
            radius = g('rx') * scaleOf / k;
          } else {
            const rx = tag === 'circle' ? g('r') : g('rx');
            const ry = tag === 'circle' ? g('r') : g('ry');
            [x0, y0] = point(m, g('cx') - rx, g('cy') - ry);
            [x1, y1] = point(m, g('cx') + rx, g('cy') + ry);
          }
          const w = Math.abs(x1 - x0), h = Math.abs(y1 - y0);
          if (tag !== 'rect') radius = Math.min(w, h) / 2;
          items.push({...base, kind: 'box', x: Math.min(x0, x1), y: Math.min(y0, y1), w, h, radius,
                      fill, stroke, stroke_width: sw});
        }
        continue;
      }
      if (tag === 'svg' || tag === 'br') continue;
      const r = rel(el.getBoundingClientRect());
      if (tag === 'img') {
        items.push({...base, kind: 'image', src: el.currentSrc || el.src, ...r, fit: cs.objectFit,
                    natural_w: el.naturalWidth, natural_h: el.naturalHeight});
        continue;
      }
      const bg = color(cs.backgroundColor);
      const sides = ['Top', 'Right', 'Bottom', 'Left'].map((s) => ({
        side: s, w: cs['border' + s + 'Style'] === 'none' || cs['border' + s + 'Style'] === 'hidden'
          ? 0 : px(cs['border' + s + 'Width']), c: cs['border' + s + 'Color']}));
      const uniform = sides.every((s) => s.w === sides[0].w && s.c === sides[0].c);
      const stroke = uniform && sides[0].w > 0 ? color(sides[0].c) : null;
      let radiusRaw = cs.borderTopLeftRadius.split(' ')[0];
      const radius = radiusRaw.endsWith('%') ? px(radiusRaw) / 100 * Math.min(r.w, r.h) : px(radiusRaw);
      const hasBox = !!(bg || stroke);
      if (!uniform) {
        // Single-side rules (accent bars, dividers) become thin filled bars.
        for (const s of sides) {
          const c = color(s.c);
          if (!s.w || !c) continue;
          const bar = {Top: [r.x, r.y, r.w, s.w], Bottom: [r.x, r.y + r.h - s.w, r.w, s.w],
                       Left: [r.x, r.y, s.w, r.h], Right: [r.x + r.w - s.w, r.y, s.w, r.h]}[s.side];
          items.push({...base, key: base.key + '#border-' + s.side.toLowerCase(), kind: 'box',
                      x: bar[0], y: bar[1], w: bar[2], h: bar[3], radius: 0, fill: c, stroke: null});
        }
      }
      const directText = [...el.childNodes].some((n) => n.nodeType === 3 && n.data.trim());
      const allInline = [...el.children].every(
        (c) => c.tagName === 'BR' || ['inline', 'contents'].includes(getComputedStyle(c).display));
      // An element holding only inline content owns all of its text; one that mixes
      // text with block children owns just its direct text.
      const ownsText = allInline ? el.textContent.trim().length > 0 : directText;
      let text = null;
      if (ownsText) {
        if (allInline) el.querySelectorAll('*').forEach((d) => consumed.add(d));
        const runs = runsOf(el, !allInline);
        if (runs.length) {
          const range = document.createRange();
          range.selectNodeContents(el);
          const tr = rel(range.getBoundingClientRect());
          const pad = ['Top', 'Right', 'Bottom', 'Left'].map((s) => px(cs['padding' + s]) + px(cs['border' + s + 'Width']));
          const align = {center: 'center', right: 'right', end: 'right', '-webkit-center': 'center'}[cs.textAlign] || 'left';
          const contentTop = r.y + pad[0], contentBottom = r.y + r.h - pad[2];
          const gapTop = tr.y - contentTop, gapBottom = contentBottom - (tr.y + tr.h);
          const valign = Math.abs(gapTop - gapBottom) <= 4 ? 'middle' : gapTop < gapBottom ? 'top' : 'bottom';
          text = {runs, align, wrap: !/nowrap|pre$/.test(cs.whiteSpace), pad, valign,
                  tx: r.x + pad[3], tw: Math.max(r.w - pad[1] - pad[3], tr.w), ty: tr.y, th: tr.h};
        }
      }
      if (hasBox) {
        const box = {...base, key: base.key + '#box', kind: 'box', ...r, radius, fill: bg, stroke,
                     stroke_width: stroke ? sides[0].w : 0};
        if (text) Object.assign(box, {runs: text.runs, align: text.align, wrap: text.wrap,
                                      padding: text.pad, valign: text.valign});
        items.push(box);
      } else if (text) {
        items.push({...base, key: base.key + '#text', kind: 'text', runs: text.runs, align: text.align,
                    wrap: text.wrap, x: text.tx, y: text.ty, w: text.tw, h: text.th});
      }
    }
    const theme = {};
    const scs = getComputedStyle(slide);
    for (const name of THEME_VARS) {
      const c = color(scs.getPropertyValue(name).trim());
      if (c) theme[name] = c;
    }
    const rr = rel(root.getBoundingClientRect());
    return {signature: signature(root), ...rr, items, theme};
  }
  function run() {
    const out = {diagrams: {}, unsupported: []};
    try {
      const unsupported = new Set();
      for (const root of document.querySelectorAll('[' + ATTR + ']')) {
        out.diagrams[root.getAttribute(ATTR)] = measure(root, unsupported);
      }
      out.unsupported = [...unsupported];
    } catch (e) {
      out.error = String(e && e.message || e);
    }
    const node = document.createElement('script');
    node.type = 'application/json';
    node.id = 'colloquium-diagram-layout';
    node.textContent = JSON.stringify(out).replace(/</g, '\\u003c');
    document.body.appendChild(node);
  }
  const start = () => (document.fonts ? document.fonts.ready : Promise.resolve()).then(run);
  if (document.readyState === 'complete') start(); else window.addEventListener('load', start);
})();
"""
